"""Export → Excel → Original File: one worksheet with the
document's sections, numbered items and tables in reading order, built from
the hierarchy the layout already decided (original_layout.py) — nothing is
re-derived here.

- A section opens with its printed heading.
- Fields sit under Item # | Field Name | Extracted Value. An item with
  sub-fields opens with its number and name on a row of their own; each
  sub-field follows on its own row, its number cell blank. An item that is
  one field is one row. Form boxes sharing a number show it once. A
  printed date range ("01/01/2025 - 12/31/2025") is followed by its Start
  Date and End Date rows; the printed text is kept in a comment. Field
  names are the form's own, without instructions ("Enter the …").
- A form's own tables (a table of contents, its amendments) are tables:
  the item's number and caption above them, their printed columns, a
  choice box as Selected / Not Selected. Rows the text extractor read
  from a form's boxes or printed wording aren't repeated beside them.
- A table keeps its printed headers, column order and row order. Document
  metadata (Contractor Name, Contract Number…) appears once above it,
  never in its rows: the block the document prints there, or for a table
  without one, its contract's from the nearest block before it. A column
  naming one contract or item for every row is written there too, once,
  instead of in each row. Its source pages are named.
- A value that isn't a confirmed source value (unreadable, missing
  extraction, unverified OCR, conflicting candidates) carries a cell
  comment saying so; a blank printed cell stays blank. Nothing is corrected
  or filled in.

No cell is merged, so every value can be edited, sorted and filtered. Every
value is written as text exactly as printed ("07/01/2025", "$33,747,047.53"),
so Excel doesn't turn it into a date or a number, and text that looks like a
formula is never executed."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.source_structure.original_layout import BLANK, CONFIRMED, EXACT, EXTRACTED, QUALITY_LABELS, SOURCE_LOCATION_LABELS
from app.staging.custom_export import safe_cell

FIELD_COLUMNS = ["Item #", "Field Name", "Extracted Value"]
_MIN_WIDTH, _MAX_WIDTH = 10, 60
_TABLE_FIRST_COLUMN = 16
# Rows whose text runs across the empty cells beside it, unwrapped.
_BANNERS = {"section", "subsection", "table_title", "context"}

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_SECTION_FILL = PatternFill("solid", fgColor="DCE6F1")
_HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
_HEADER_EDGE = Border(bottom=Side(style="thin", color="A6A6A6"))
_TOP = Alignment(vertical="top", wrap_text=True)
_BANNER = Alignment(vertical="top", wrap_text=False)
# Sub-fields sit under their item, set in by one level.
_INDENTED = Alignment(vertical="top", wrap_text=True, indent=1)
_CHILDREN = {"subfield", "part"}
_AUTHOR = "Dataomatix"


@dataclass
class Row:
    """One worksheet row: what it is, its cells, and a note per cell."""

    kind: str  # section | subsection | header | item | field | part | table_title | context | table_header | table_row | gap
    values: list[Any] = field(default_factory=list)
    notes: dict[int, str] = field(default_factory=dict)


def field_value(item: dict[str, Any]) -> str:
    """The value now in effect: a reviewer's correction, else what was read
    (a ticked form box reads "Selected")."""

    if item.get("choice") and not item.get("structure_review"):
        return item["choice"]
    return item.get("effective_value") or item.get("extracted_value") or item.get("document_value") or ""


def _location_note(item: dict[str, Any]) -> str | None:
    """Where the value can be traced, said apart from its quality."""

    location = item.get("source_location")
    if not location or location == EXACT:
        return None
    page = f" (page {item['page']})" if item.get("page") else ""
    return f"Source: {SOURCE_LOCATION_LABELS.get(location, location)}{page}"


def _field_note(item: dict[str, Any]) -> str | None:
    code = item.get("quality") or CONFIRMED
    if code in (CONFIRMED, BLANK, EXTRACTED):
        return _location_note(item)
    note = QUALITY_LABELS.get(code, code)
    if item.get("ocr_confidence") is not None:
        note += f" (OCR confidence {item['ocr_confidence']:.2f})"
    elif item.get("ocr_contested"):
        note += " (the OCR passes disagreed)"
    others = [m.get("value") or m.get("original_value") for m in item.get("other_metadata") or []]
    others = [v for v in others if v]
    if others:
        note += "; other candidates: " + "; ".join(dict.fromkeys(others))
    if item.get("document_value") and item.get("document_value") != field_value(item):
        note += f"; printed: {item['document_value']}"
    location = _location_note(item)
    return f"{note}; {location}" if location else note


def _composite_note(item: dict[str, Any]) -> str:
    parts = " and ".join(part["label"] for part in item["parts"])
    return f"Printed as {field_value(item)}; split into {parts}"


def _table_context(item: dict[str, Any]) -> list[str]:
    """What a table's rows share, once: its inherited contract block and
    any column holding one value for every row."""

    entries = []
    if item.get("context_source") == "inherited":
        # As the document printed the label ("Legal Entity"), not the concept.
        entries = [f"{entry.get('printed_label') or entry['label']}: {entry['value']}" for entry in item.get("context") or []]
    return entries + [f"{shared['label']}: {shared['value']}" for shared in item.get("shared_columns") or []]


def _pages(pages: list[int]) -> str:
    if not pages:
        return ""
    return f"Source: page {pages[0]}" if len(pages) == 1 else f"Source: pages {pages[0]}–{pages[-1]}"


def original_rows(layout: dict[str, Any]) -> list[Row]:
    """The worksheet, row by row, in the document's reading order."""

    rows: list[Row] = []

    def gap() -> None:
        if rows and rows[-1].kind != "gap":
            rows.append(Row("gap"))

    sections = list(layout["sections"])
    if layout.get("unplaced"):
        sections.append({"title": "Extracted fields without a printed position", "level": 1, "items": layout["unplaced"]})
    for section in sections:
        # A section whose only items repeat the form's boxes isn't listed.
        if not section["items"] or all(i.get("form_duplicate") for i in section["items"]):
            continue
        gap()
        title = section.get("title") or ""
        if title:
            rows.append(Row("section" if (section.get("level") or 1) == 1 else "subsection", [title]))
        in_fields = False
        previous_item: str | None = None
        for item in section["items"]:
            if item.get("form_duplicate"):
                # The text extractor's second reading of a form's boxes or
                # wording: the form's own boxes are written instead.
                continue
            if item["kind"] == "field":
                if not in_fields:
                    rows.append(Row("header", list(FIELD_COLUMNS)))
                    in_fields = True
                parent = item.get("item")
                value = field_value(item)
                name = item.get("field_name") or item["label"]
                first = parent is not None and parent["id"] != previous_item
                kind = "field"
                if parent is not None and parent.get("has_subfields") and parent.get("name"):
                    # The item's number and name on a row of their own, once.
                    if first:
                        rows.append(Row("item", [parent["number"], parent["name"], ""]))
                    values, kind = ["", name, value], "subfield"
                elif parent is not None:
                    values = [parent["number"] if first else "", name, value]
                    if not first:
                        kind = "subfield"
                else:
                    values = ["", name, value]
                previous_item = parent["id"] if parent is not None else None
                note = _field_note(item)
                parts = item.get("parts") or []
                if parts:
                    # The printed text once, as a note; its parts carry the values.
                    values[2] = ""
                    note = "; ".join(filter(None, [_composite_note(item), note]))
                rows.append(Row(kind, values, {2: note} if note else {}))
                for part in parts:
                    rows.append(Row("part", ["", part["label"], part["value"]]))
                continue
            in_fields = False
            previous_item = None
            gap()
            name = item.get("title") if item.get("title") and item.get("title") != title else "Table"
            if item.get("item"):
                # A form's numbered table: its item number in the Item # column.
                rows.append(Row("table_title", [item["item"]["number"], item["item"]["name"] or name, _pages(item.get("pages") or [])]))
            else:
                rows.append(Row("table_title", [name, _pages(item.get("pages") or [])]))
            # Metadata the document prints right above the table is already
            # there as fields; a table without its own gets its contract's.
            context = _table_context(item)
            if context:
                # One line, so it reads across the table without widening its first column.
                rows.append(Row("context", ["   ·   ".join(context)]))
            width = len(item["headers"])
            shared = {entry["column"] for entry in item.get("shared_columns") or []}
            kept = [column for column in range(width) if column not in shared]
            rows.append(Row("table_header", [item["headers"][column] for column in kept]))
            for row in item["rows"]:
                cells = (list(row.get("reviewed_cells") or row["cells"]) + [""] * width)[:width]
                notes = {i["column"]: i["message"] for i in row.get("cell_issues") or [] if i["column"] < width}
                for flag in row.get("flags") or []:
                    if flag.get("column") is not None and flag["column"] < width:
                        notes[flag["column"]] = "; ".join(filter(None, [notes.get(flag["column"]), flag["message"]]))
                rows.append(Row(
                    "table_row",
                    [cells[column] for column in kept],
                    {kept.index(column): note for column, note in notes.items() if column in kept},
                ))
            gap()
    while rows and rows[-1].kind == "gap":
        rows.pop()
    return rows


def csv_rows(layout: dict[str, Any]) -> list[list[str]]:
    """The same rows as plain values (a CSV has no comments or styles)."""

    return [[str(v) if v is not None else "" for v in row.values] for row in original_rows(layout)]


def _text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    return safe_cell(_ILLEGAL.sub(" ", str(value)))


_STYLE = {
    "item": (Font(bold=True), None, None),
    "section": (Font(bold=True, size=12), _SECTION_FILL, None),
    "subsection": (Font(bold=True, size=11), None, None),
    "header": (Font(bold=True), _HEADER_FILL, _HEADER_EDGE),
    "table_header": (Font(bold=True), _HEADER_FILL, _HEADER_EDGE),
    "table_title": (Font(bold=True), None, None),
    "context": (Font(italic=True, color="44546A"), None, None),
}


def build_original_structure_xlsx(layout: dict[str, Any]) -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Original File"
    widths: dict[int, int] = {}
    for row_number, row in enumerate(original_rows(layout), start=1):
        if row.kind not in _BANNERS:
            for index, value in enumerate(row.values):
                longest = max((len(line) for line in str(value or "").splitlines()), default=0)
                # Column A is also the Item # column: a table's long first
                # cells wrap there rather than widen it for every field row.
                cap = _TABLE_FIRST_COLUMN if index == 0 and row.kind in ("table_header", "table_row") else _MAX_WIDTH
                widths[index] = max(widths.get(index, _MIN_WIDTH), min(cap, longest + 2))
        font, fill, border = _STYLE.get(row.kind, (None, None, None))
        # Header and section rows are painted across the field columns so
        # they read as one band; each value still sits in its own cell.
        span = max(len(row.values), len(FIELD_COLUMNS) if fill else 0)
        for column in range(1, span + 1):
            value = row.values[column - 1] if column <= len(row.values) else None
            text = _text(value)
            cell = ws.cell(row=row_number, column=column)
            if text is not None:
                cell.value = text
                cell.data_type = "s"  # never a formula, never re-typed
            cell.number_format = "@"
            cell.alignment = _BANNER if row.kind in _BANNERS else _INDENTED if row.kind in _CHILDREN and column == 2 else _TOP
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
            note = row.notes.get(column - 1)
            if note:
                cell.comment = Comment(note, _AUTHOR, width=280, height=90)
    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index + 1)].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

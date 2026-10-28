"""Custom export: a reviewer-defined table built for importing elsewhere.

The browser sends a mapping — output columns, each a label and one or more
source fields, by id — never values. Values come from the document's
staging workbook, exactly as Export All reads them.

* A source with a record id is one value (a document-level field: an
  invoice number, a passport's surname, a Key Field's value).
* A source without one is a column of a repeating dataset (invoice lines):
  the file gets one row per record of that dataset, and every one-value
  column repeats on each row. Only one such dataset per export, so rows are
  unambiguous; with none, the file is a single row.
* Several sources in one column are combined, empty values skipped, joined
  by the column's separator ("Given Names" + "Surname" → "ANNA ERIKSSON").

The file is import-ready: headers in the first row (unless turned off), one
record per row, no title rows, no styling that importers trip over.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pydantic import BaseModel, Field

from app.staging.export import find_dataset
from app.staging.models import StagingDataset, StagingWorkbook
from app.staging.selection_export import ExportFormat, SelectionError, is_exportable

_MEDIA = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "json": "application/json",
}
_MAX_ROWS = 200_000


class ColumnSource(BaseModel):
    dataset_id: str = Field(min_length=1, max_length=120)
    field: str = Field(min_length=1, max_length=200)
    # Set: that record's value. Unset: a column of the repeating dataset.
    record_id: str | None = Field(default=None, max_length=300)


class CustomColumn(BaseModel):
    label: str = Field(min_length=1, max_length=120)
    sources: list[ColumnSource] = Field(min_length=1, max_length=20)
    separator: str = Field(default=" ", max_length=10)


class CustomExportManifest(BaseModel):
    format: ExportFormat
    columns: list[CustomColumn] = Field(min_length=1, max_length=500)
    include_header: bool = True
    file_name: str | None = Field(default=None, max_length=120)


def _value_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _dataset(workbook: StagingWorkbook, dataset_id: str) -> StagingDataset:
    dataset = find_dataset(workbook, dataset_id)
    if dataset is None:
        raise SelectionError(f"Unknown dataset {dataset_id!r}.")
    return dataset


def _check_field(dataset: StagingDataset, field: str) -> None:
    if not any(column.canonical_field == field and is_exportable(column) for column in dataset.columns):
        raise SelectionError(f"{dataset.display_name} has no exportable field {field!r}.")


def build_rows(workbook: StagingWorkbook, manifest: CustomExportManifest) -> tuple[list[str], list[list[object]]]:
    """(headers, rows) — each cell a plain value (combined cells are text)."""

    labels = [column.label.strip() for column in manifest.columns]
    if any(not label for label in labels):
        raise SelectionError("Every column needs a label.")
    duplicates = sorted({label for label in labels if labels.count(label) > 1})
    if duplicates:
        raise SelectionError(f"Column labels must be unique for importing: {', '.join(duplicates)}.")

    row_datasets = {
        source.dataset_id for column in manifest.columns for source in column.sources if source.record_id is None
    }
    if len(row_datasets) > 1:
        raise SelectionError("Table columns must all come from one table, so each row is one record.")
    row_dataset = _dataset(workbook, next(iter(row_datasets))) if row_datasets else None
    if row_dataset is not None and row_dataset.cardinality != "repeating":
        raise SelectionError(f"{row_dataset.display_name} is not a table; pick its fields as single values.")
    if row_dataset is not None and len(row_dataset.records) > _MAX_ROWS:
        raise SelectionError("Too many rows to export at once.")

    # One-value sources: resolved once, repeated on every row.
    fixed: dict[tuple[str, str, str], object] = {}
    for column in manifest.columns:
        for source in column.sources:
            dataset = _dataset(workbook, source.dataset_id)
            _check_field(dataset, source.field)
            if source.record_id is None:
                continue
            record = next((r for r in dataset.records if r.record_id == source.record_id), None)
            if record is None:
                raise SelectionError(f"{dataset.display_name} has no record {source.record_id!r}.")
            cell = record.cells.get(source.field)
            fixed[(source.dataset_id, source.record_id, source.field)] = cell.value if cell else None

    records = row_dataset.records if row_dataset is not None else [None]
    rows: list[list[object]] = []
    for record in records:
        row: list[object] = []
        for column in manifest.columns:
            values = []
            for source in column.sources:
                if source.record_id is None:
                    cell = record.cells.get(source.field) if record is not None else None
                    values.append(cell.value if cell else None)
                else:
                    values.append(fixed[(source.dataset_id, source.record_id, source.field)])
            if len(values) == 1:
                row.append(values[0])
            else:
                row.append(column.separator.join(_value_text(v) for v in values if _value_text(v).strip()))
        rows.append(row)
    return labels, rows


def _csv(headers: list[str], rows: list[list[object]], include_header: bool) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    if include_header:
        writer.writerow(headers)
    for row in rows:
        writer.writerow([_value_text(value) for value in row])
    return buffer.getvalue()


def _json(headers: list[str], rows: list[list[object]]) -> str:
    return json.dumps([dict(zip(headers, row)) for row in rows], ensure_ascii=False, indent=1, default=str)


def _xlsx(headers: list[str], rows: list[list[object]], include_header: bool) -> bytes:
    """A plain sheet: headers in row 1 (frozen, bold), values from row 2 — no
    title rows, so importers read it as a table."""

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    if include_header:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    for row in rows:
        # Strings that look like formulas are written as text, never evaluated.
        sheet.append([
            ("'" + value) if isinstance(value, str) and value[:1] in "=+-@" and not re.fullmatch(r"[-+]?\d[\d,.]*", value)
            else value
            for value in row
        ])
    for index, header in enumerate(headers, start=1):
        sample = [len(header)] + [len(_value_text(row[index - 1])) for row in rows[:200]]
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, max(sample) + 2))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _filename(workbook: StagingWorkbook, manifest: CustomExportManifest) -> str:
    stem = manifest.file_name or f"{(workbook.document_filename or 'export').rsplit('.', 1)[0]}_custom"
    stem = re.sub(r"[^\w\- .]+", "_", stem).strip(" .") or "export"
    stem = re.sub(r"\.(xlsx|csv|json)$", "", stem, flags=re.IGNORECASE)
    return f"{stem}.{manifest.format}"


def build_custom_export(
    workbook: StagingWorkbook, manifest: CustomExportManifest
) -> tuple[str | bytes, str, str]:
    """(content, media type, filename) for the manifest's format."""

    headers, rows = build_rows(workbook, manifest)
    if manifest.format == "csv":
        content: str | bytes = _csv(headers, rows, manifest.include_header)
    elif manifest.format == "json":
        content = _json(headers, rows)
    else:
        content = _xlsx(headers, rows, manifest.include_header)
    return content, _MEDIA[manifest.format], _filename(workbook, manifest)


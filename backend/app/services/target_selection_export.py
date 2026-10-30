"""Export Selected for extracted fields and table rows (Source Verification).

The request names records by id only — a field's normalized key, a table
row by its table key and position in the **stored** table — and the values
are read from the persisted extraction results, never from the browser.
Unselected records are never written. Cells are formula-safe.
"""

from __future__ import annotations

import csv
import io
import json
from typing import Any, Literal

from pydantic import BaseModel, Field

from app.schemas.document_target import ExtractTargetsResponse
from app.staging.custom_export import safe_cell

ExportFormat = Literal["xlsx", "csv", "json"]


class SelectedRow(BaseModel):
    table: str
    index: int = Field(ge=0)


class TargetSelection(BaseModel):
    format: ExportFormat = "xlsx"
    fields: list[str] = Field(default_factory=list, max_length=5000)
    rows: list[SelectedRow] = Field(default_factory=list, max_length=50000)


class SelectionError(ValueError):
    pass


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def selected_records(result: ExtractTargetsResponse, selection: TargetSelection) -> tuple[list[dict], list[dict]]:
    by_key = {scalar.normalized_key: scalar for scalar in result.scalars}
    unknown = [key for key in selection.fields if key not in by_key]
    if unknown:
        raise SelectionError(f"Unknown fields: {', '.join(unknown[:5])}")
    fields = [
        {
            "field": by_key[key].target,
            "key": key,
            "value": _text(by_key[key].value),
            "page": by_key[key].evidence.page_number or by_key[key].page,
            "section": by_key[key].evidence.section or "",
            "validation": by_key[key].validation_status,
            "review_status": by_key[key].review_status or "",
        }
        for key in dict.fromkeys(selection.fields)  # selection order, each once
    ]
    tables = {table.target: table for table in result.tables}
    grouped: dict[str, list[int]] = {}
    for row in selection.rows:
        table = tables.get(row.table)
        if table is None or row.index >= len(table.rows):
            raise SelectionError(f"Unknown table row: {row.table} #{row.index}")
        grouped.setdefault(row.table, [])
        if row.index not in grouped[row.table]:
            grouped[row.table].append(row.index)
    table_rows = [
        {
            "table": name,
            "columns": list(tables[name].columns),
            "pages": list(tables[name].pages),
            # Stored order, whatever order the rows were ticked in.
            "rows": [{"index": i, "cells": {c: _text(tables[name].rows[i].get(c)) for c in tables[name].columns}} for i in sorted(indexes)],
        }
        for name, indexes in grouped.items()
    ]
    return fields, table_rows


FIELD_HEADERS = ["Field", "Value", "Page", "Section", "Validation", "Review status"]


def build_selection_export(result: ExtractTargetsResponse, selection: TargetSelection, stem: str) -> tuple[bytes, str, str]:
    fields, tables = selected_records(result, selection)
    if not fields and not tables:
        raise SelectionError("Nothing is selected.")
    field_rows = [[f["field"], f["value"], f["page"], f["section"], f["validation"], f["review_status"]] for f in fields]

    if selection.format == "json":
        body = json.dumps({"fields": fields, "tables": tables}, ensure_ascii=False, indent=2).encode()
        return body, "application/json", f"{stem}_selected.json"

    if selection.format == "csv":
        # One table: fields as rows; table rows cell by cell (Table | Row | Column | Value).
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["Record", *FIELD_HEADERS, "Table", "Row", "Column"])
        for row in field_rows:
            writer.writerow([safe_cell(v) for v in ["Field", *row, "", "", ""]])
        for table in tables:
            for row in table["rows"]:
                for column in table["columns"]:
                    writer.writerow([safe_cell(v) for v in ["Table row", "", row["cells"][column], "", "", "", "", table["table"], row["index"] + 1, column]])
        return ("﻿" + buffer.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", f"{stem}_selected.csv"

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Selected fields"
    sheet.append([safe_cell(h) for h in FIELD_HEADERS])
    for row in field_rows:
        sheet.append([safe_cell(v) for v in row])
    for number, table in enumerate(tables, start=1):
        rows_sheet = book.create_sheet(f"Table {number}"[:31])
        rows_sheet.append([safe_cell(f"{table['table']} (pages {', '.join(map(str, table['pages']))})")])
        rows_sheet.append([safe_cell(c) for c in table["columns"]])
        for row in table["rows"]:
            rows_sheet.append([safe_cell(row["cells"][c]) for c in table["columns"]])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{stem}_selected.xlsx"

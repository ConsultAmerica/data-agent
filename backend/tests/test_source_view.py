"""Increment 6B: structured source views (app/source_view)."""

from __future__ import annotations

import json
import re
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.source_view import service
from app.xml_records.parse import parse_tree

client = TestClient(app)


def xlsx_with_formula(cached: str | None = "1250") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pricing"
    sheet.append(["CLIN", "Qty", "Rate", "Amount"])
    sheet.append(["0001", 5, 250, "=B2*C2"])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Merged note"
    sheet.freeze_panes = "B2"
    workbook.create_sheet("Notes")["A1"] = "Second sheet"
    buffer = BytesIO()
    workbook.save(buffer)
    if cached is None:
        return buffer.getvalue()
    # Excel stores a formula's last result next to it (<f>…</f><v>…</v>);
    # openpyxl can't write one, so add it the way Excel would.
    source = zipfile.ZipFile(BytesIO(buffer.getvalue()))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"(<c r=\"D2\"[^>]*>)<f>B2\*C2</f>(<v\s*/>|<v></v>)?", rb"\1<f>B2*C2</f><v>" + cached.encode() + rb"</v>", data)
            target.writestr(item, data)
    return out.getvalue()


ODS = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
 <office:body><office:spreadsheet>
  <table:table table:name="Rates">
   <table:table-row><table:table-cell><text:p>Item</text:p></table:table-cell><table:table-cell><text:p>Total</text:p></table:table-cell></table:table-row>
   <table:table-row table:number-rows-repeated="3"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
   <table:table-row>
    <table:table-cell table:number-columns-repeated="2"/>
    <table:table-cell table:formula="of:=[.A1]*2" office:value="84"><text:p>84</text:p></table:table-cell>
    <table:table-cell table:number-columns-spanned="2"><text:p>Wide</text:p></table:table-cell>
   </table:table-row>
  </table:table>
 </office:spreadsheet></office:body>
</office:document-content>"""


def make_ods() -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        archive.writestr("content.xml", ODS, zipfile.ZIP_DEFLATED)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = DocxDocument()
    document.core_properties.title = "Statement of Work"
    document.add_heading("1. Scope", level=1)
    document.add_paragraph("The contractor shall deliver <b>services</b>.")
    document.add_paragraph("First deliverable", style="List Bullet")
    document.add_heading("1.1 Schedule", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "CLIN", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "0001", "100.00"
    paragraph = document.add_paragraph("See ")
    rel = paragraph.part.relate_to("https://example.com/terms", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "the terms"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


XML = b"""<?xml version="1.0"?>
<inv:Invoice xmlns:inv="urn:example:invoice" xmlns:cac="urn:example:common" currency="USD">
  <inv:Number>INV-7</inv:Number>
  <cac:Line id="1"><cac:Amount>10.00</cac:Amount></cac:Line>
  <cac:Line id="2"><cac:Amount>12.50</cac:Amount></cac:Line>
</inv:Invoice>"""


def stored(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


# ── viewer routing ───────────────────────────────────────────────────


@pytest.mark.parametrize(
    "name, data, viewer",
    [
        ("a.pdf", b"%PDF-1.4 x", "pdf"),
        ("a.png", b"\x89PNG\r\n\x1a\nxxxx", "image"),
        ("a.xlsx", None, "spreadsheet"),
        ("a.xlsm", None, "spreadsheet"),
        ("a.ods", None, "spreadsheet"),
        ("a.csv", b"a,b\n1,2\n", "delimited"),
        ("a.tsv", b"a\tb\n", "delimited"),
        ("a.json", b"{}", "json"),
        ("a.xml", b"<a/>", "xml"),
        ("a.docx", None, "word"),
        ("a.txt", b"hello", "text"),
        ("a.doc", b"\xd0\xcf\x11\xe0rest", "text"),
    ],
)
def test_viewer_is_chosen_from_the_verified_type(tmp_path, name, data, viewer) -> None:
    if data is None:
        data = {"a.xlsx": xlsx_with_formula(None), "a.xlsm": xlsx_with_formula(None), "a.ods": make_ods(), "a.docx": make_docx()}[name]
    resolved = service.resolve_viewer(name, stored(tmp_path, name, data))
    assert resolved["viewer"] == viewer
    assert resolved["verified"] is True
    if name == "a.doc":
        assert "original layout isn't reproduced" in resolved["note"]


def test_a_file_whose_content_contradicts_its_type_is_not_previewed(tmp_path) -> None:
    resolved = service.resolve_viewer("fake.xlsx", stored(tmp_path, "fake.xlsx", b"%PDF-1.4 not a workbook"))
    assert resolved["viewer"] == "unsupported" and resolved["verified"] is False
    assert service.resolve_viewer("a.exe", stored(tmp_path, "a.exe", b"MZ"))["viewer"] == "unsupported"


# ── spreadsheets ─────────────────────────────────────────────────────


def test_xlsx_formula_and_cached_value_are_separate_and_nothing_is_calculated(tmp_path) -> None:
    path = stored(tmp_path, "p.xlsx", xlsx_with_formula("1250"))
    summary = service.spreadsheet_summary(path, "xlsx")
    assert [s["name"] for s in summary["sheets"]] == ["Pricing", "Notes"]
    assert summary["sheets"][0]["frozen"] == "B2"
    sheet = service.spreadsheet_sheet(path, "xlsx", 0)
    cells = {c["ref"]: c for c in sheet["cells"]}
    assert cells["D2"]["formula"] == "=B2*C2" and cells["D2"]["value"] == "1250"
    assert cells["A2"] == {"ref": "A2", "row": 2, "column": 1, "formula": None, "value": "0001", "type": "s"}
    assert sheet["merged"] == ["A4:B4"]

    # Never calculated in the file: no value is made up.
    uncached = service.spreadsheet_sheet(stored(tmp_path, "u.xlsx", xlsx_with_formula(None)), "xlsx", 0)
    d2 = next(c for c in uncached["cells"] if c["ref"] == "D2")
    assert d2["formula"] == "=B2*C2" and d2["value"] is None


def test_xlsm_says_macros_are_never_loaded(tmp_path) -> None:
    summary = service.spreadsheet_summary(stored(tmp_path, "m.xlsm", xlsx_with_formula(None)), "xlsm")
    assert "never loaded or run" in summary["macros_note"]


def test_ods_cells_keep_true_positions_formulas_and_spans(tmp_path) -> None:
    path = stored(tmp_path, "r.ods", make_ods())
    sheet = service.spreadsheet_sheet(path, "ods", 0)
    cells = {c["ref"]: c for c in sheet["cells"]}
    assert set(cells) == {"A1", "B1", "C5", "D5"}
    assert cells["C5"]["formula"] == "of:=[.A1]*2" and cells["C5"]["value"] == "84"
    assert sheet["merged"] == ["D5:E5"]
    assert sheet["max_row"] == 5 and sheet["max_column"] == 5


def test_column_letters() -> None:
    assert [service.column_letter(n) for n in (1, 26, 27, 52, 703)] == ["A", "Z", "AA", "AZ", "AAA"]


# ── delimited ────────────────────────────────────────────────────────


def test_delimited_view_keeps_text_and_reports_its_dialect(tmp_path) -> None:
    path = stored(tmp_path, "r.csv", 'CLIN;Description;Amount\n0001;"Labor; senior";1.234,50\n0002;Travel;99\n'.encode("cp1252"))
    view = service.delimited_view(path)
    assert view["delimiter"] == "semicolon"
    assert view["rows"][1] == ["0001", "Labor; senior", "1.234,50"]
    assert view["encoding_guessed"] is False  # plain ASCII is valid UTF-8
    ragged = service.delimited_view(stored(tmp_path, "g.csv", "a,b\n1,2,3\nZürich,x\n".encode("cp1252")))
    assert ragged["encoding_guessed"] is True
    assert any("different numbers of fields" in w for w in ragged["warnings"])
    assert any("encoding" in w for w in ragged["warnings"])


# ── XML ──────────────────────────────────────────────────────────────


def test_xml_tree_keeps_namespaces_and_uses_the_extraction_paths(tmp_path) -> None:
    view = service.xml_view(stored(tmp_path, "i.xml", XML))
    root = view["root"]
    assert (root["name"], root["prefix"], root["namespace"]) == ("Invoice", "inv", "urn:example:invoice")
    assert {(n["prefix"], n["uri"]) for n in view["namespaces"]} == {("inv", "urn:example:invoice"), ("cac", "urn:example:common")}
    assert [c["path"] for c in root["children"]] == ["/Invoice/Number", "/Invoice/Line[1]", "/Invoice/Line[2]"]
    assert root["children"][2]["attributes"][0] == {"name": "id", "namespace": None, "value": "2", "path": "/Invoice/Line[2]/@id"}
    assert root["children"][2]["children"][0]["text"] == "12.50"

    # The same paths XML extraction records as evidence (xml_records._paths).
    from app.xml_records.parse import _paths

    tree = parse_tree(XML)
    assert sorted(_paths(tree).values()) == sorted(
        ["/Invoice", "/Invoice/Number", "/Invoice/Line[1]", "/Invoice/Line[2]", "/Invoice/Line[1]/Amount", "/Invoice/Line[2]/Amount"]
    )


def test_xml_with_entities_is_refused(tmp_path) -> None:
    evil = b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>'
    with pytest.raises(service.SourceViewError):
        service.xml_view(stored(tmp_path, "e.xml", evil))


# ── Word ─────────────────────────────────────────────────────────────


def test_word_blocks_are_structured_data_never_html(tmp_path) -> None:
    view = service.word_view(stored(tmp_path, "s.docx", make_docx()))
    kinds = [(b["type"], b.get("level")) for b in view["blocks"]]
    assert kinds[:3] == [("heading", 1), ("paragraph", None), ("list_item", 0)]
    assert ("heading", 2) in kinds and ("table", None) in kinds
    paragraph = view["blocks"][1]
    assert paragraph["text"] == "The contractor shall deliver <b>services</b>."  # literal text, not markup
    table = next(b for b in view["blocks"] if b["type"] == "table")
    assert table["rows"] == [["CLIN", "Amount"], ["0001", "100.00"]]
    assert table["section"] == ["1. Scope", "1.1 Schedule"]
    link = next(p for b in view["blocks"] for p in b.get("parts", []) if p.get("href"))
    assert link == {"text": "the terms", "href": "https://example.com/terms", "anchor": None}
    assert view["metadata"]["title"] == "Statement of Work"


def test_oversized_sources_fall_back(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(service, "MAX_SOURCE_BYTES", 10)
    with pytest.raises(service.SourceViewError, match="larger than"):
        service.json_view(stored(tmp_path, "big.json", b'{"a": "0123456789"}'))


# ── API ──────────────────────────────────────────────────────────────


def upload(name: str, data: bytes, mime: str) -> str:
    response = client.post("/api/documents/upload", files={"file": (name, data, mime)})
    assert response.status_code == 201, response.text
    return response.json()["document_id"]


def test_source_view_api_routes_each_format() -> None:
    marker = uuid4().hex[:8]
    xlsx_id = upload("p.xlsx", xlsx_with_formula("1250"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    view = client.get(f"/api/documents/{xlsx_id}/source-view").json()
    assert view["viewer"] == "spreadsheet" and view["verified"] is True
    sheet = client.get(f"/api/documents/{xlsx_id}/source-view/sheets/0").json()
    assert any(c["ref"] == "D2" and c["formula"] == "=B2*C2" for c in sheet["cells"])
    assert client.get(f"/api/documents/{xlsx_id}/source-view/sheets/9").status_code == 422

    csv_id = upload("r.csv", f"id,name\n007,{marker}\n".encode(), "text/csv")
    view = client.get(f"/api/documents/{csv_id}/source-view").json()
    assert view["viewer"] == "delimited" and view["rows"][1] == ["007", marker]
    assert client.get(f"/api/documents/{csv_id}/source-view/sheets/0").status_code == 422

    json_id = upload("d.json", json.dumps({"a": [marker]}).encode(), "application/json")
    assert client.get(f"/api/documents/{json_id}/source-view").json()["data"] == {"a": [marker]}

    xml_id = upload("i.xml", XML.replace(b"INV-7", marker.encode()), "application/xml")
    assert client.get(f"/api/documents/{xml_id}/source-view").json()["root"]["children"][0]["text"] == marker

    docx_id = upload("s.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert client.get(f"/api/documents/{docx_id}/source-view").json()["viewer"] == "word"

    assert client.get(f"/api/documents/{uuid4()}/source-view").status_code == 404


# ── regression: spreadsheet and CSV documents open their workspace ──


def test_table_shapes_accept_both_stored_forms() -> None:
    from app.services.table_shapes import as_table, page_tables

    assert as_table({"headers": ["a"], "rows": [["1"]]}) == {"headers": ["a"], "rows": [["1"]]}
    assert as_table([["CLIN", "Amount"], ["0001", "5"]]) == {"headers": ["CLIN", "Amount"], "rows": [["0001", "5"]]}
    assert as_table("junk") is None and as_table([]) is None
    assert len(list(page_tables([{"headers": [], "rows": []}, [["a"]], None]))) == 2


@pytest.mark.parametrize(
    "name, data, mime",
    [
        ("book.xlsx", None, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("two.csv", None, "text/csv"),
    ],
)
def test_spreadsheet_and_csv_documents_build_their_staging_workbook(name, data, mime) -> None:
    # Native tables are stored as lists of rows; the generic profile's
    # key/value scanner used to crash on them (500 on staging-workbook).
    marker = uuid4().hex[:8]
    if name.endswith(".xlsx"):
        workbook = Workbook()
        workbook.active.append(["Label", "Value"])
        workbook.active.append(["Reference", marker])
        buffer = BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
    else:
        data = f"Label,Value\nReference,{marker}\n".encode()
    document_id = upload(name, data, mime)
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}).status_code == 200
    response = client.get(f"/api/documents/{document_id}/staging-workbook")
    assert response.status_code == 200, response.text[:300]

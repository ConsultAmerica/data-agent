"""Oracle Template Manager: template parsing, mapping, validation and
exact-header export.

Template headers below are test data, not verified Oracle metadata: the
platform takes Oracle names only from uploaded templates, so the tests use
templates the way a user would upload them.
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.oracle_templates.mapping import (
    ColumnMapping,
    MappingSpec,
    ParentLink,
    SourceRef,
    apply_transform,
    build_rows,
    catalogue,
    import_ready,
    suggest,
    suggest_links,
    validate,
    write_csv,
)
from app.oracle_templates.parser import TemplateError, parse_template
from app.staging.models import (
    CellProvenance,
    ExtractionOutcomeModel,
    ProcessingMetadata,
    ProfileDescriptor,
    QaSummary,
    StagingCell,
    StagingColumn,
    StagingDataset,
    StagingRecord,
    StagingWorkbook,
)

client = TestClient(app)

REGRESSION = Path(__file__).resolve().parents[2] / "reference" / "regression"
NAVY = REGRESSION / "SF1442 Award_N4019223D2803.pdf"

HEADER_TECHNICAL = "BatchId,ExtSource,ExtKey,ContractNumber,ContractName,EffectiveStartDate,CurrencyCode,EstimatedAmount"
HEADER_LABELS = "Batch,Source,*Key,*Number,Name,Start Date,Currency,Amount"
LINE_TECHNICAL = "HeaderExtSource,HeaderExtKey,LineNumber,ItemDescription,Quantity,UnitPrice,Amount,LineAction"


# ---------------------------------------------------------------- synthetic workbook


def _cell(field: str, label: str, value, page: int | None = 1, status: str | None = "Verified") -> StagingCell:
    provenance = (
        CellProvenance(source_document_id="d1", source_filename="c.pdf", source_type="pdf", source_page=page,
                       highlight_text=str(value), anchor_text=label.upper())
        if value is not None else None
    )
    return StagingCell(canonical_field=field, display_label=label, value=value, provenance=provenance,
                       review_status=status if value is not None else "Missing")


def _dataset(dataset_id: str, name: str, cardinality: str, columns: list[tuple[str, str, str]], rows: list[list]) -> StagingDataset:
    cols = [StagingColumn(canonical_field=field, key=key, display_label=label) for key, field, label in columns]
    return StagingDataset(
        dataset_id=dataset_id, display_name=name, cardinality=cardinality, columns=cols,
        records=[
            StagingRecord(record_id=f"{dataset_id}:{i}",
                          cells={c.canonical_field: _cell(c.canonical_field, c.display_label, v) for c, v in zip(cols, row)})
            for i, row in enumerate(rows)
        ],
    )


def _workbook(lines: list[list] | None = None) -> StagingWorkbook:
    summary = _dataset("contract_summary", "Contract Summary", "single",
                       [("contract_name", "contract.contract_name", "Contract Name"),
                        ("currency", "contract.currency", "Currency")],
                       [["Base Operations, \"BOS\" Support", "USD"]])
    details = _dataset("contract_details", "Contract Details", "repeating",
                       [("label", "contract.detail.label", "Field"), ("value", "contract.detail.value", "Value"),
                        ("field_id", "contract.detail.field_id", "Field Id")],
                       [["2. CONTRACT NUMBER", "FA300224C0008", "contract.contract_number"],
                        ["20. AMOUNT", "USD 24,131,140.28", "contract.award_amount"],
                        ["28. AWARD DATE", "21 Apr 2023", "contract.award_date"]])
    line_items = _dataset("line_items", "Line Items", "repeating",
                          [("item_number", "contract.line.item_number", "Item Number"),
                           ("description", "contract.line.description", "Description"),
                           ("quantity", "contract.line.quantity", "Quantity"),
                           ("unit_price", "contract.line.unit_price", "Unit Price"),
                           ("amount", "contract.line.amount", "Amount")],
                          lines if lines is not None else [
                              ["0001", "ALL SERVICES", "10", "USD 142,035.17", "USD 1,420,351.70"],
                              ["0002", "Travel, as directed", "1", "Not to Exceed", "USD 30,000.00"],
                          ])
    return StagingWorkbook(
        document_id="d1", document_filename="c.pdf",
        profile=ProfileDescriptor(profile_id="contract_v3", profile_version=1, display_name="Contract", description="",
                                  document_families=["government_contract"], export_capabilities=[]),
        outcome=ExtractionOutcomeModel(status="populated", title="", message=""),
        datasets=[summary, details, line_items], qa_summary=QaSummary(), processing_metadata=ProcessingMetadata(),
    )


def _columns(header: str, labels: str | None = None) -> list[dict]:
    parsed = parse_template("Contract.csv", (header + "\n" + (labels + "\n" if labels else "")).encode())
    return [vars(c) for c in parsed.columns]


def _pair(field_id: str) -> SourceRef:
    return SourceRef(kind="field", dataset_id="contract_details", field="contract.detail.value", shape="pair",
                     pair_by="field_id", pair_key=field_id)


def _header_spec() -> MappingSpec:
    return MappingSpec(
        mappings=[
            ColumnMapping(column="ExtSource", source=SourceRef(kind="constant", constant="DATAOMATIX")),
            ColumnMapping(column="ExtKey", source=_pair("contract.contract_number")),
            ColumnMapping(column="ContractNumber", source=_pair("contract.contract_number")),
            ColumnMapping(column="ContractName", source=SourceRef(kind="field", dataset_id="contract_summary",
                                                                  field="contract.contract_name", shape="single")),
            ColumnMapping(column="EffectiveStartDate", source=_pair("contract.award_date"), transform="date_iso"),
            ColumnMapping(column="CurrencyCode", source=SourceRef(kind="field", dataset_id="contract_summary",
                                                                  field="contract.currency", shape="single")),
            ColumnMapping(column="EstimatedAmount", source=_pair("contract.award_amount"), transform="number"),
        ],
        key_columns=["ExtSource", "ExtKey"],
    )


def _line_spec(parent_id: int | None = None) -> MappingSpec:
    def column(field: str) -> SourceRef:
        return SourceRef(kind="field", dataset_id="line_items", field=field, shape="column")

    return MappingSpec(
        row_dataset="line_items",
        mappings=[
            ColumnMapping(column="HeaderExtSource", source=SourceRef(kind="constant", constant="DATAOMATIX")),
            ColumnMapping(column="HeaderExtKey", source=_pair("contract.contract_number")),
            ColumnMapping(column="LineNumber", source=column("contract.line.item_number")),
            ColumnMapping(column="ItemDescription", source=column("contract.line.description")),
            ColumnMapping(column="Quantity", source=column("contract.line.quantity")),
            ColumnMapping(column="UnitPrice", source=column("contract.line.unit_price"), transform="number"),
            ColumnMapping(column="Amount", source=column("contract.line.amount"), transform="number"),
            ColumnMapping(column="LineAction", source=SourceRef(kind="constant", constant="ADD")),
        ],
        key_columns=["HeaderExtKey", "LineNumber"],
        parent=ParentLink(profile_id=parent_id, links=[("HeaderExtSource", "ExtSource"), ("HeaderExtKey", "ExtKey")])
        if parent_id is not None else None,
    )


# ---------------------------------------------------------------- parser


def test_two_header_rows_give_technical_names_labels_and_required_marks() -> None:
    parsed = parse_template("Contract.csv", f"{HEADER_TECHNICAL}\n{HEADER_LABELS}\n".encode())
    assert parsed.object_name == "Contract"
    assert (parsed.technical_row, parsed.label_row) == (0, 1)
    assert [c.technical_name for c in parsed.columns] == HEADER_TECHNICAL.split(",")
    assert [c.display_label for c in parsed.columns][:4] == ["Batch", "Source", "Key", "Number"]
    assert [c.required for c in parsed.columns][:4] == [None, None, True, True]


def test_labels_first_layout_and_semicolon_delimiter() -> None:
    parsed = parse_template("Contract_Line.csv", b"Number;Name\nContractNumber;ContractName\n")
    assert parsed.object_name == "Contract Line"
    assert (parsed.technical_row, parsed.label_row) == (1, 0)
    assert [(c.technical_name, c.display_label) for c in parsed.columns] == [
        ("ContractNumber", "Number"), ("ContractName", "Name"),
    ]


def test_technical_only_template_reports_that_required_is_unknown() -> None:
    parsed = parse_template("Contract_Party.csv", f"{LINE_TECHNICAL}\n".encode())
    assert parsed.label_row is None
    assert all(c.display_label is None and c.required is None for c in parsed.columns)
    assert any("required" in warning for warning in parsed.warnings)


def test_sample_data_rows_are_not_taken_for_labels() -> None:
    parsed = parse_template("Contract.csv", b"ContractNumber,ContractName\nFA300224C0008,Base Operations Support\n")
    assert parsed.label_row is None


def test_label_only_template_keeps_its_header_strings_and_warns() -> None:
    parsed = parse_template("Contract.csv", b"Number,Primary Party\n")
    assert [c.technical_name for c in parsed.columns] == ["Number", "Primary Party"]
    assert any("display labels" in warning for warning in parsed.warnings)


def test_workbook_template_skips_instruction_sheets_and_never_runs_macros() -> None:
    book = Workbook()
    book.active.title = "Instructions"
    book.active.append(["Fill in the Contract sheet."])
    sheet = book.create_sheet("Contract")
    sheet.append(HEADER_TECHNICAL.split(","))
    sheet.append(HEADER_LABELS.split(","))
    buffer = io.BytesIO()
    book.save(buffer)
    parsed = parse_template("ContractImportTemplate.xlsm", buffer.getvalue())
    assert parsed.sheet_name == "Contract" and parsed.object_name == "Contract"
    assert [c.technical_name for c in parsed.columns] == HEADER_TECHNICAL.split(",")
    assert parsed.warnings[0].startswith("Macro-enabled workbook")


@pytest.mark.parametrize("content, message", [
    (b"", "empty"),
    (b"ContractNumber,ContractNumber\n", "Duplicate"),
    (b"ContractNumber,,ContractName\n", "empty"),
])
def test_unusable_templates_are_refused(content: bytes, message: str) -> None:
    with pytest.raises(TemplateError, match=message):
        parse_template("Contract.csv", content)


# ---------------------------------------------------------------- rows, transforms, csv


def test_header_object_is_one_row_in_template_order_with_exact_headers() -> None:
    built = build_rows(_workbook(), _columns(HEADER_TECHNICAL, HEADER_LABELS), _header_spec())
    assert built.headers == HEADER_TECHNICAL.split(",")
    assert [cell.value for cell in built.rows[0]] == [
        "", "DATAOMATIX", "FA300224C0008", "FA300224C0008", 'Base Operations, "BOS" Support', "2023-04-21", "USD", "24131140.28",
    ]
    content = write_csv(built.headers, [[c.value for c in row] for row in built.rows])
    assert not content.startswith(b"\xef\xbb\xbf")
    lines = content.decode("utf-8").splitlines()
    assert lines[0] == HEADER_TECHNICAL  # exact header text and order, no display labels
    assert list(csv.reader(io.StringIO(lines[1])))[0][4] == 'Base Operations, "BOS" Support'


def test_line_object_gives_one_row_per_record_and_keeps_leading_zeros() -> None:
    built = build_rows(_workbook(), _columns(LINE_TECHNICAL), _line_spec())
    assert [[c.value for c in row][:3] for row in built.rows] == [
        ["DATAOMATIX", "FA300224C0008", "0001"], ["DATAOMATIX", "FA300224C0008", "0002"],
    ]
    assert built.rows[0][6].value == "1420351.70"
    # "Not to Exceed" is not a price: kept as printed and reported, never guessed.
    assert built.rows[1][5].value == "Not to Exceed" and "not a number" in built.rows[1][5].problem


@pytest.mark.parametrize("value, transform, expected", [
    ("21 Apr 2023", "date_iso", "2023-04-21"),
    ("04/15/2025", "date_iso", "2025-04-15"),
    ("Sep 29, 2023", "date_iso", "2023-09-29"),
    ("USD 1,000.50", "number", "1000.50"),
    ("(250.00)", "number", "-250.00"),
    ("00042", "as_is", "00042"),
    ("  Base   Ops ", "trim", "Base Ops"),
])
def test_transforms(value: str, transform: str, expected: str) -> None:
    assert apply_transform(value, transform) == (expected, None)


# ---------------------------------------------------------------- validation


def test_validation_uses_only_what_the_template_declares() -> None:
    columns = _columns(HEADER_TECHNICAL, HEADER_LABELS)  # ExtKey and ContractNumber marked required
    spec = _header_spec()
    issues = validate(columns, spec, build_rows(_workbook(), columns, spec))
    assert import_ready(issues)
    assert {i.code for i in issues if i.severity == "info"} >= {"unmapped", "lookups_not_checked"}

    spec.mappings = [m for m in spec.mappings if m.column != "ContractNumber"]
    issues = validate(columns, spec, build_rows(_workbook(), columns, spec))
    assert not import_ready(issues)
    assert any(i.code == "required_unmapped" and i.column == "ContractNumber" for i in issues)

    # Without required marks in the template, nothing is treated as required.
    bare = _columns(HEADER_TECHNICAL)
    assert import_ready(validate(bare, spec, build_rows(_workbook(), bare, spec)))


def test_duplicate_and_missing_keys_and_data_type_errors() -> None:
    columns = _columns(LINE_TECHNICAL)
    workbook = _workbook(lines=[["0001", "A", "1", "10", "10"], ["0001", "B", "1", "x", "20"], ["", "C", "1", "5", "5"]])
    issues = validate(columns, _line_spec(), build_rows(workbook, columns, _line_spec()))
    codes = {(i.code, i.row) for i in issues}
    assert ("duplicate_key", 2) in codes and ("missing_identifier", 3) in codes and ("data_type", 2) in codes


def test_child_rows_must_match_a_parent_key() -> None:
    header_columns = _columns(HEADER_TECHNICAL, HEADER_LABELS)
    parent = build_rows(_workbook(), header_columns, _header_spec())
    columns = _columns(LINE_TECHNICAL)
    spec = _line_spec(parent_id=1)
    spec.mappings = [m for m in spec.mappings if m.column != "HeaderExtSource"] + [
        ColumnMapping(column="HeaderExtSource", source=SourceRef(kind="constant", constant="OTHER")),
    ]
    issues = validate(columns, spec, build_rows(_workbook(), columns, spec), parent)
    assert [i.code for i in issues if i.code == "orphan_row"] == ["orphan_row", "orphan_row"]
    assert suggest_links(LINE_TECHNICAL.split(","), HEADER_TECHNICAL.split(",")) == [
        ("HeaderExtSource", "ExtSource"), ("HeaderExtKey", "ExtKey"),
    ]


def test_mapping_to_a_column_the_template_lacks_is_a_compatibility_error() -> None:
    columns = _columns(HEADER_TECHNICAL)
    spec = MappingSpec(mappings=[ColumnMapping(column="NotInTemplate", source=SourceRef(kind="constant", constant="x"))])
    assert any(i.code == "template_mismatch" for i in validate(columns, spec, build_rows(_workbook(), columns, spec)))


# ---------------------------------------------------------------- suggestions


def test_suggestions_explain_their_method_and_pick_the_row_table() -> None:
    fields = catalogue(_workbook())
    header, row = suggest(_columns(HEADER_TECHNICAL, HEADER_LABELS), fields, object_name="Contract")
    by_column = {s.column: s for s in header}
    assert row is None
    assert by_column["ContractNumber"].source.pair_key == "contract.contract_number"
    assert by_column["ContractNumber"].method == "exact name"
    assert by_column["CurrencyCode"].method == "label"  # the Oracle label "Currency"

    lines, row = suggest(_columns(LINE_TECHNICAL), fields, object_name="Contract Line")
    by_column = {s.column: s for s in lines}
    assert row == "line_items"
    assert by_column["Amount"].source.field == "contract.line.amount"  # the line's, not the contract's
    assert by_column["LineNumber"].source.field == "contract.line.item_number"


def test_previously_approved_mapping_is_suggested_first() -> None:
    fields = catalogue(_workbook())
    learned = {"extkey": _pair("contract.contract_number")}
    suggestions, _ = suggest(_columns(HEADER_TECHNICAL), fields, learned, object_name="Contract")
    ext_key = next(s for s in suggestions if s.column == "ExtKey")
    assert (ext_key.method, ext_key.confidence) == ("previously approved", 0.95)


def test_catalogue_carries_document_labels_and_evidence() -> None:
    fields = {f.label: f for f in catalogue(_workbook())}
    assert fields["2. CONTRACT NUMBER"].sample == "FA300224C0008"
    assert fields["2. CONTRACT NUMBER"].ref.pair_by == "field_id"
    assert fields["Currency"].page == 1 and fields["Currency"].document_label == "CURRENCY"


# ---------------------------------------------------------------- API: end to end on a real contract


@pytest.fixture(scope="module")
def navy_document() -> str:
    if not NAVY.exists():
        pytest.skip("SF1442 regression contract not present")
    from tests.test_export_consistency import _process

    return _process(NAVY.name, NAVY.read_bytes(), "application/pdf")


def _upload(name: str, content: bytes, **form):
    return client.post("/api/oracle/templates", files={"file": (name, content, "text/csv")}, data=form)


def _csv_rows(response) -> list[list[str]]:
    assert not response.content.startswith(b"\xef\xbb\xbf")
    return list(csv.reader(io.StringIO(response.content.decode("utf-8"))))


def test_template_manager_and_exact_header_export_end_to_end(navy_document: str) -> None:
    document_id = navy_document
    header_bytes = f"{HEADER_TECHNICAL}\n{HEADER_LABELS}\n".encode()

    # 1-2. Upload: technical names, labels and order detected; the file is kept byte for byte.
    uploaded = _upload("Contract.csv", header_bytes, environment="Test pod")
    assert uploaded.status_code == 201, uploaded.text
    template = uploaded.json()
    assert [c["technical_name"] for c in template["columns"]] == HEADER_TECHNICAL.split(",")
    assert template["columns"][3]["display_label"] == "Number" and template["has_display_labels"]
    assert client.get(f"/api/oracle/templates/{template['id']}/original").content == header_bytes
    again = _upload("Contract.csv", header_bytes).json()
    assert (again["id"], again["reused"]) == (template["id"], True)  # identical bytes: same template
    second = _upload("Contract.csv", header_bytes + b"\n").json()
    assert second["object_name"] == "Contract" and second["version"] == template["version"] + 1

    # 3. Suggestions from the document's extracted fields.
    suggested = client.post(f"/api/documents/{document_id}/oracle/templates/{template['id']}/suggest", json={})
    assert suggested.status_code == 200, suggested.text
    by_column = {s["column"]: s for s in suggested.json()["suggestions"]}
    assert by_column["ContractNumber"]["source"]["kind"] == "field"
    assert by_column["ContractNumber"]["method"] in {"exact name", "label", "previously approved"}

    # 4-5. Change a mapping by hand and save it.
    spec = {
        "row_dataset": None,
        "mappings": [
            {"column": s["column"], "source": s["source"], "method": s["method"], "confidence": s["confidence"]}
            for s in suggested.json()["suggestions"] if s["column"] in {"ContractNumber", "CurrencyCode"}
        ] + [
            {"column": "ExtSource", "source": {"kind": "constant", "constant": "DATAOMATIX"}, "method": "manual"},
            {"column": "ExtKey", "source": by_column["ContractNumber"]["source"], "method": "manual"},
        ],
        "key_columns": ["ExtSource", "ExtKey"],
    }
    created = client.post(f"/api/oracle/templates/{template['id']}/profiles",
                          json={"name": "Navy contract header", "staging_profile_id": "contract_v3", "spec": spec})
    assert created.status_code == 201, created.text
    profile = created.json()
    assert profile["version"] == 1 and len(profile["history"]) == 1

    # 6-7. Generate the CSV: exact headers, template order, one row.
    exported = client.post(f"/api/documents/{document_id}/oracle/templates/{template['id']}/export", json={"spec": profile["spec"]})
    assert exported.status_code == 200, exported.text
    rows = _csv_rows(exported)
    assert rows[0] == HEADER_TECHNICAL.split(",")
    assert len(rows) == 2
    number = rows[1][HEADER_TECHNICAL.split(",").index("ContractNumber")]
    assert number.replace("-", "") == "N4019223D2803"
    assert rows[1][1] == "DATAOMATIX" and rows[1][2] == number
    assert exported.headers["x-import-ready"] == "true"

    # 8. Change the mapping, save a new version, regenerate.
    changed = profile["spec"]
    changed["mappings"] = [m for m in changed["mappings"] if m["column"] != "ExtSource"] + [
        {"column": "ExtSource", "source": {"kind": "constant", "constant": "LEGACY"}, "transform": "as_is"},
    ]
    saved = client.put(f"/api/oracle/profiles/{profile['id']}",
                       json={"name": "Navy contract header", "spec": changed})
    assert saved.status_code == 200 and saved.json()["version"] == 2 and len(saved.json()["history"]) == 2
    regenerated = client.post(f"/api/documents/{document_id}/oracle/templates/{template['id']}/export",
                              json={"spec": saved.json()["spec"]})
    assert _csv_rows(regenerated)[1][1] == "LEGACY"

    # 9. Unmapped extracted data stays available (and listed) — nothing is dropped.
    preview = client.post(f"/api/documents/{document_id}/oracle/templates/{template['id']}/preview",
                          json={"spec": saved.json()["spec"]}).json()
    assert preview["import_ready"] and preview["total_rows"] == 1
    assert preview["display_labels"][3] == "Number"
    assert len(preview["unmapped_fields"]) > 10
    assert client.get(f"/api/documents/{document_id}/staging-workbook").status_code == 200

    # Lines, linked to the header, bundled in one ZIP named after the templates.
    line_template = _upload("Contract_Line.csv", f"{LINE_TECHNICAL}\n".encode()).json()
    links = client.post(f"/api/documents/{document_id}/oracle/templates/{line_template['id']}/suggest",
                        json={"parent_profile_id": profile["id"]}).json()
    assert links["row_dataset"] and [tuple(link) for link in links["parent_links"]] == [
        ("HeaderExtSource", "ExtSource"), ("HeaderExtKey", "ExtKey"),
    ]
    line_by_column = {s["column"]: s for s in links["suggestions"]}
    line_spec = {
        "row_dataset": links["row_dataset"],
        "mappings": [
            {"column": "HeaderExtSource", "source": {"kind": "constant", "constant": "LEGACY"}},
            {"column": "HeaderExtKey", "source": by_column["ContractNumber"]["source"]},
            {"column": "LineNumber", "source": line_by_column["LineNumber"]["source"]},
            {"column": "LineAction", "source": {"kind": "constant", "constant": "ADD"}},
        ],
        "key_columns": ["HeaderExtKey", "LineNumber"],
        "parent": {"profile_id": profile["id"], "links": links["parent_links"]},
    }
    line_profile = client.post(f"/api/oracle/templates/{line_template['id']}/profiles",
                               json={"name": "Navy lines", "spec": line_spec}).json()
    bundle = client.post(f"/api/documents/{document_id}/oracle/export-bundle",
                         json={"profile_ids": [profile["id"], line_profile["id"]], "force": True})
    assert bundle.status_code == 200, bundle.text
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
        assert sorted(archive.namelist()) == ["Contract.csv", "Contract_Line.csv"]
        lines = list(csv.reader(io.StringIO(archive.read("Contract_Line.csv").decode("utf-8"))))
    assert lines[0] == LINE_TECHNICAL.split(",")
    assert len(lines) > 2 and all(row[1] == number for row in lines[1:])


def test_export_is_refused_while_validation_finds_errors(navy_document: str) -> None:
    template = _upload("Strict_Contract.csv", b"ContractNumber,ContractName\n*Number,*Name\n").json()
    spec = {"mappings": [{"column": "ContractNumber", "source": {"kind": "constant", "constant": "X"}}]}
    refused = client.post(f"/api/documents/{navy_document}/oracle/templates/{template['id']}/export", json={"spec": spec})
    assert refused.status_code == 409
    forced = client.post(f"/api/documents/{navy_document}/oracle/templates/{template['id']}/export",
                         json={"spec": spec, "force": True})
    assert forced.status_code == 200 and forced.headers["x-import-ready"] == "false"


def test_bad_template_upload_is_rejected_with_a_reason() -> None:
    response = _upload("Contract.csv", b"ContractNumber,ContractNumber\n")
    assert response.status_code == 422 and "Duplicate" in response.json()["detail"]
    response = client.post("/api/oracle/templates", files={"file": ("contract.pdf", b"%PDF-1.4", "application/pdf")})
    assert response.status_code == 422


def test_a_header_repeated_for_a_lookups_id_and_name_is_kept_in_place() -> None:
    # Oracle's contract line export prints a lookup as two columns under one
    # label: the Supplier's id, then its name.
    # A header row of display labels (Oracle's own export of contract lines).
    parsed = parse_template("ContractLine.csv", b"Number,LineId,Supplier,Supplier,Status,Status,Unit Price\n")
    names = [c.technical_name for c in parsed.columns]
    assert names == ["Number", "LineId", "Supplier", "Supplier (2)", "Status", "Status (2)", "Unit Price"]
    assert [c.header for c in parsed.columns] == [None, None, None, "Supplier", None, "Status", None]
    assert any("Repeated header names (Status, Supplier)" in w for w in parsed.warnings)


def test_an_export_prints_a_repeated_header_as_the_template_does() -> None:
    columns = [vars(c) for c in parse_template("ContractLine.csv", b"Supplier,Supplier,Contract Name\n").columns]
    spec = MappingSpec(mappings=[
        ColumnMapping(column="Supplier (2)", source=SourceRef(kind="constant", constant="ACME Corp")),
        ColumnMapping(column="Contract Name", source=SourceRef(kind="constant", constant="Base Operations")),
    ])
    built = build_rows(_workbook(), columns, spec)
    assert built.headers == ["Supplier", "Supplier", "Contract Name"]
    rows = list(csv.reader(io.StringIO(write_csv(built.headers, [[c.value for c in r] for r in built.rows]).decode())))
    assert rows[0] == ["Supplier", "Supplier", "Contract Name"]
    assert rows[1] == ["", "ACME Corp", "Base Operations"]

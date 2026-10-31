"""Phase B (7.3.3): the Oracle Fusion mapping workspace — target catalog,
verified schema registration, metadata discovery, transformations,
attribute validation, contract isolation, profile lifecycle and versions,
preview, clause sources, tenant isolation.

Every Oracle attribute below is test data in an uploaded template, as a
user would register it; nothing here is presented as real Oracle metadata.
"""

from __future__ import annotations

import csv
import io
import json
import secrets
from contextlib import contextmanager
from unittest.mock import patch
from uuid import uuid4

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.config import get_settings
from app.main import app
from app.oracle_templates import discovery
from app.oracle_templates.catalog import status_of
from app.oracle_templates.mapping import (
    ColumnMapping,
    MappingSpec,
    ParentLink,
    SourceRef,
    apply_mapping,
    build_rows,
    clause_records,
    column_states,
    contract_numbers,
    suggest,
    synonyms,
    validate,
)
from app.oracle_templates.parser import TemplateError, apply_definitions, parse_attribute_definitions
from tests.test_oracle_templates import _columns, _dataset, _header_spec, _pair, _workbook

client = TestClient(app)

HEADER = "ExtSource,ExtKey,ContractNumber,ContractName,EffectiveStartDate,CurrencyCode,EstimatedAmount"
LABELS = "Source,*Key,*Number,Name,Start Date,Currency,Amount"
DEFINITIONS = (
    "Attribute Name,Data Type,Required,Lookup,Description\n"
    "ContractNumber,VARCHAR2(20),Y,,The contract's number\n"
    'EstimatedAmount,"NUMBER(12,2)",N,,\n'
    "CurrencyCode,VARCHAR2(3),N,USD;EUR,\n"
    "EffectiveStartDate,DATE,N,,\n"
)


def _template_bytes(header: str = HEADER, labels: str = LABELS) -> bytes:
    return f"{header}\n{labels}\n".encode()


def _register(name: str = "Contract", content: bytes | None = None, **form) -> dict:
    response = client.post(
        "/api/oracle/templates",
        # A unique trailing line: identical bytes would reuse an earlier upload.
        files={"file": (f"{name}.csv", content or _template_bytes() + f"\n#{uuid4().hex}".encode(), "text/csv")},
        data={"object_name": f"{name} {uuid4().hex[:6]}", **form},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _verified(**catalog) -> dict:
    template = _register(application="Enterprise Contracts", integration_method="fbdi", oracle_release="25A", **catalog)
    defined = client.post(f"/api/oracle/templates/{template['id']}/attributes",
                          files={"file": ("definitions.csv", DEFINITIONS.encode(), "text/csv")})
    assert defined.status_code == 200, defined.text
    verified = client.post(f"/api/oracle/templates/{template['id']}/verify")
    assert verified.status_code == 200, verified.text
    return verified.json()


def _document() -> str:
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Contract Number: FA300224C0008", fontsize=10)
    response = client.post("/api/documents/upload", files={"file": (f"map-{uuid4()}.pdf", pdf.tobytes(), "application/pdf")})
    assert response.status_code == 201, response.text
    return response.json()["document_id"]


@contextmanager
def _workbook_for_documents(workbook):
    with patch("app.api.oracle_templates.get_staging_workbook", return_value=workbook):
        yield


# ---------------------------------------------------------------- 1-3. catalog and schemas


def test_a_registered_template_is_a_draft_until_verified_then_listed_under_its_application() -> None:
    template = _register(application="Enterprise Contracts", integration_method="fbdi", oracle_release="25A",
                         object_code="OKC_K_HEADERS", user_keys="ExtSource,ExtKey")
    catalog = template["catalog"]
    assert (catalog["status"], catalog["source"]) == ("draft", "uploaded_template")
    assert catalog["user_keys"] == ["ExtSource", "ExtKey"] and template["required_columns"] == ["ExtKey", "ContractNumber"]

    verified = client.post(f"/api/oracle/templates/{template['id']}/verify").json()
    assert verified["catalog"]["status"] == "verified" and verified["catalog"]["verified_by"]

    listing = client.get("/api/oracle/catalog").json()
    assert listing["configured"] and listing["application_names"][0] == "Enterprise Contracts"
    group = next(a for a in listing["applications"] if a["application"] == "Enterprise Contracts")
    assert any(o["template"]["id"] == template["id"] for o in group["objects"])
    assert "fbdi" in listing["integration_methods"] and "lookup" in listing["transformations"]

    # Changing what the target is withdraws the verification.
    changed = client.put(f"/api/oracle/templates/{template['id']}/catalog",
                         json={"application": "Enterprise Contracts", "integration_method": "fbdi", "oracle_release": "25B"})
    assert changed.json()["catalog"]["status"] == "draft"


def test_unsupported_methods_unknown_applications_and_malformed_files_are_refused() -> None:
    # Import Management is the Sales and Service mechanism, not Enterprise Contracts'.
    wrong = _register(application="Enterprise Contracts", integration_method="import_management")
    assert wrong["catalog"]["status"] == "unsupported"
    assert client.post(f"/api/oracle/templates/{wrong['id']}/verify").status_code == 422
    unknown = client.post("/api/oracle/templates", files={"file": ("x.csv", _template_bytes(), "text/csv")},
                          data={"application": "Oracle Everything", "integration_method": "fbdi"})
    assert unknown.status_code == 422
    keys = client.post("/api/oracle/templates", files={"file": ("y.csv", _template_bytes() + b"\n#k", "text/csv")},
                       data={"application": "Procurement", "integration_method": "fbdi", "user_keys": "NotAColumn"})
    assert keys.status_code == 422 and "User keys" in keys.text
    # Not a template: no header row.
    assert client.post("/api/oracle/templates", files={"file": ("z.csv", b"\n\n", "text/csv")}).status_code == 422
    # A definitions file needs an attribute column and a property column.
    template = _register()
    bad = client.post(f"/api/oracle/templates/{template['id']}/attributes",
                      files={"file": ("notes.csv", b"Note,Comment\nhello,world\n", "text/csv")})
    assert bad.status_code == 422
    unrelated = client.post(f"/api/oracle/templates/{template['id']}/attributes",
                            files={"file": ("d.csv", b"Attribute Name,Data Type\nSomethingElse,DATE\n", "text/csv")})
    assert unrelated.status_code == 422


def test_attribute_definitions_give_types_lengths_precision_and_lookups() -> None:
    definitions = parse_attribute_definitions("d.csv", DEFINITIONS.encode())
    assert definitions["ContractNumber"] == {"data_type": "string", "declared_type": "VARCHAR2(20)", "max_length": 20,
                                             "required": True, "description": "The contract's number"}
    assert (definitions["EstimatedAmount"]["precision"], definitions["EstimatedAmount"]["scale"]) == (12, 2)
    assert definitions["CurrencyCode"]["lookup_values"] == ["USD", "EUR"]
    columns, warnings = apply_definitions(_columns(HEADER, LABELS), definitions)
    by_name = {c["technical_name"]: c for c in columns}
    assert by_name["EffectiveStartDate"]["data_type"] == "date" and not warnings
    # A template's own required mark is never overridden.
    assert by_name["ContractNumber"]["required"] is True

    with pytest.raises(TemplateError):
        parse_attribute_definitions("d.csv", b"")


def test_a_workbook_template_documents_its_own_attributes() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contract"
    sheet.append(HEADER.split(","))
    sheet.append(LABELS.split(","))
    notes = workbook.create_sheet("Attributes")
    for row in csv.reader(DEFINITIONS.strip().splitlines()):
        notes.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = client.post("/api/oracle/templates",
                           files={"file": ("Contract.xlsx", buffer.getvalue(), "application/octet-stream")},
                           data={"object_name": f"Contract WB {uuid4().hex[:6]}"})
    assert response.status_code == 201, response.text
    by_name = {c["technical_name"]: c for c in response.json()["columns"]}
    assert by_name["ContractNumber"]["max_length"] == 20 and by_name["CurrencyCode"]["lookup_values"] == ["USD", "EUR"]


# ---------------------------------------------------------------- 4. suggestions


def test_suggestions_say_required_need_approval_and_check_data_types(tmp_path) -> None:
    columns, _ = apply_definitions(_columns(HEADER, LABELS), parse_attribute_definitions("d.csv", DEFINITIONS.encode()))
    from app.oracle_templates.mapping import catalogue

    suggestions, _ = suggest(columns, catalogue(_workbook()))
    by_column = {s.column: s for s in suggestions}
    number = by_column["ContractNumber"]
    assert number.required and number.needs_approval
    amount = by_column["EstimatedAmount"]
    assert "fits its number type" in amount.reason

    # Aliases are configurable: a file of extra groups adds names.
    aliases = tmp_path / "aliases.json"
    aliases.write_text(json.dumps({"groups": [["contract name", "program title"]]}), encoding="utf-8")
    settings = get_settings().model_copy(update={"oracle_mapping_aliases_path": str(aliases)})
    synonyms.cache_clear()
    try:
        with patch("app.core.config.get_settings", return_value=settings):
            assert any({"contract name", "program title"} <= group for group in synonyms())
    finally:
        synonyms.cache_clear()


# ---------------------------------------------------------------- 6-8. validation and transformations


def _mapped(column: str, source: SourceRef, transform: str = "as_is", **options) -> ColumnMapping:
    return ColumnMapping(column=column, source=source, transform=transform, options=options)


def test_required_type_length_precision_and_lookup_are_checked_against_the_attributes() -> None:
    columns, _ = apply_definitions(_columns(HEADER, LABELS), parse_attribute_definitions("d.csv", DEFINITIONS.encode()))
    workbook = _workbook()
    summary = SourceRef(kind="field", dataset_id="contract_summary", field="contract.contract_name", shape="single")
    spec = MappingSpec(mappings=[
        _mapped("ContractNumber", summary),  # 30 characters: too long for VARCHAR2(20)
        _mapped("EstimatedAmount", _pair("contract.award_amount")),  # not converted: not a number
        _mapped("CurrencyCode", SourceRef(kind="constant", constant="GBP")),  # not a lookup value
        _mapped("EffectiveStartDate", _pair("contract.award_date")),  # not YYYY-MM-DD
    ])
    issues = validate(columns, spec, build_rows(workbook, columns, spec))
    codes = {(i.code, i.column) for i in issues}
    assert ("required_unmapped", "ExtKey") in codes
    assert ("length", "ContractNumber") in codes
    assert ("data_type", "EstimatedAmount") in codes and ("data_type", "EffectiveStartDate") in codes
    assert ("lookup", "CurrencyCode") in codes

    fixed = MappingSpec(mappings=[
        _mapped("ExtKey", _pair("contract.contract_number")),
        _mapped("ContractNumber", _pair("contract.contract_number")),
        _mapped("EstimatedAmount", _pair("contract.award_amount"), "currency"),
        _mapped("CurrencyCode", SourceRef(kind="constant", constant="USD")),
        _mapped("EffectiveStartDate", _pair("contract.award_date"), "date_iso"),
    ])
    fixed_issues = validate(columns, fixed, build_rows(workbook, columns, fixed))
    assert not [i for i in fixed_issues if i.severity == "error"]
    states = column_states(columns, fixed, fixed_issues)
    assert states["ContractNumber"] == "Valid" and states["ContractName"] == "Not Configured"
    # Precision: 13 digits don't fit NUMBER(12,2).
    big = validate(columns, MappingSpec(mappings=[_mapped("EstimatedAmount", SourceRef(kind="constant", constant="12345678901.25"))]),
                   build_rows(workbook, columns, MappingSpec(mappings=[_mapped("EstimatedAmount", SourceRef(kind="constant", constant="12345678901.25"))])))
    assert any(i.code == "precision" for i in big)


@pytest.mark.parametrize(
    ("value", "transform", "options", "expected", "problem"),
    [
        ("  Base  Ops ", "trim", {}, "Base Ops", False),
        ("Base Ops", "lower", {}, "base ops", False),
        ("base ops", "title", {}, "Base Ops", False),
        ("21 Apr 2023", "date_iso", {}, "2023-04-21", False),
        ("21 Apr 2023", "date_iso", {"format": "%m/%d/%Y"}, "04/21/2023", False),
        ("USD 24,131,140.28", "number", {}, "24131140.28", False),
        ("USD 24,131,140.285", "currency", {"decimals": 2}, "24131140.29", False),
        ("EUR 10.00", "currency", {"currency": "USD"}, "EUR 10.00", True),
        ("", "default", {"value": "N/A"}, "N/A", False),
        ("FFP", "lookup", {"map": {"ffp": "FIRM_FIXED_PRICE"}}, "FIRM_FIXED_PRICE", False),
        ("T&M", "lookup", {"map": {"ffp": "FIRM_FIXED_PRICE"}}, "T&M", True),
        ("Selected", "conditional", {"when": "equals", "value": "selected", "then": "Y", "else": "N"}, "Y", False),
        ("Not Selected", "conditional", {"when": "equals", "value": "selected", "then": "Y", "else": "N"}, "N", False),
        ("01/01/2025 - 12/31/2025", "split", {"part": "End Date"}, "12/31/2025", False),
        ("FA3002  338 ESS CC\nJBSA RANDOLPH, TX 78150-4300\nUNITED STATES", "split", {"part": "City"}, "JBSA RANDOLPH", False),
        ("A / B / C", "split", {"separator": "/", "index": 2}, "C", False),
        ("No range here", "split", {"part": "Start Date"}, "No range here", True),
    ],
)
def test_transformations_are_configurable_and_never_guess(value, transform, options, expected, problem) -> None:
    mapping = ColumnMapping(column="X", source=SourceRef(kind="constant", constant=value), transform=transform, options=options)
    result, found = apply_mapping(value, mapping)
    assert result == expected and bool(found) == problem


def test_concatenate_joins_its_sources_and_the_source_value_is_kept() -> None:
    columns = _columns("ContractName")
    spec = MappingSpec(mappings=[ColumnMapping(
        column="ContractName", transform="concat", options={"separator": " - "},
        source=_pair("contract.contract_number"),
        sources=[SourceRef(kind="field", dataset_id="contract_summary", field="contract.contract_name", shape="single")],
    )])
    built = build_rows(_workbook(), columns, spec)
    cell = built.rows[0][0]
    assert cell.value == 'FA300224C0008 - Base Operations, "BOS" Support'
    assert cell.source_value == "FA300224C0008"


# ---------------------------------------------------------------- 9-10. parent-child and contract isolation


def _two_contract_workbook():
    workbook = _workbook()
    workbook.datasets.append(_dataset(
        "contract_lines", "Contract Lines", "repeating",
        [("contract", "contract.line.contract_number", "Contract Number"), ("item", "contract.line.item", "Item")],
        [["FA300224C0008", "0001"], ["W912DY24C0042", "0001"]],
    ))
    return workbook


def test_records_of_different_contracts_are_never_combined() -> None:
    workbook = _two_contract_workbook()
    found = contract_numbers(workbook)
    assert found.tables == {"FA300224C0008", "W912DY24C0042"}
    columns = _columns(HEADER, LABELS)
    header = _header_spec()
    issues = validate(columns, header, build_rows(workbook, columns, header), contracts=found)
    assert any(i.code == "multiple_contracts" for i in issues)

    # Lines of two contracts under one parent record.
    line_columns = _columns("HeaderExtKey,LineNumber")
    lines = MappingSpec(
        row_dataset="contract_lines",
        mappings=[ColumnMapping(column="HeaderExtKey", source=SourceRef(kind="field", dataset_id="contract_lines",
                                                                       field="contract.line.contract_number", shape="column")),
                  ColumnMapping(column="LineNumber", source=SourceRef(kind="field", dataset_id="contract_lines",
                                                                     field="contract.line.item", shape="column"))],
        parent=ParentLink(profile_id=1, links=[("HeaderExtKey", "ExtKey")]),
    )
    parent_rows = build_rows(_workbook(), columns, header)
    line_issues = validate(line_columns, lines, build_rows(workbook, line_columns, lines), parent_rows, contracts=found)
    codes = {i.code for i in line_issues}
    assert "rows_cross_contracts" in codes and "orphan_row" in codes

    # A one-contract document: no isolation error.
    single = validate(columns, header, build_rows(_workbook(), columns, header), contracts=contract_numbers(_workbook()))
    assert not any(i.code == "multiple_contracts" for i in single)


def test_a_solicitation_number_in_a_contract_column_is_not_a_second_contract() -> None:
    workbook = _two_contract_workbook()
    workbook.datasets.append(_dataset("solicitation", "Solicitation", "single",
                                      [("solicitation", "contract.solicitation_number", "Solicitation Number")],
                                      [["W912DY24C0042"]]))
    assert contract_numbers(workbook).tables == {"FA300224C0008"}


# ---------------------------------------------------------------- 11. profiles


def _spec_for(template: dict) -> dict:
    return {"row_dataset": None, "key_columns": [], "parent": None, "mappings": [
        {"column": "ExtKey", "source": {"kind": "constant", "constant": "K1"}},
        {"column": "ContractNumber", "source": {"kind": "constant", "constant": "FA300224C0008"}},
    ]}


def test_a_profile_is_drafted_approved_changed_compared_duplicated_and_archived() -> None:
    template = _verified()
    created = client.post(f"/api/oracle/templates/{template['id']}/profiles",
                          json={"name": "FA30 Federal Contract → Oracle Enterprise Contracts", "spec": _spec_for(template),
                                "source_document_type": "Contract"})
    assert created.status_code == 201, created.text
    profile = created.json()
    assert (profile["status"], profile["version"], profile["created_by"]) == ("draft", 1, profile["updated_by"])

    # An analyst can map but not approve.
    with patch("app.core.auth.get_settings", return_value=get_settings().model_copy(update={"rbac_enforced": True})):
        denied = client.post(f"/api/oracle/profiles/{profile['id']}/activate", headers={"X-Actor-Id": "actor-analyst-default"})
    assert denied.status_code == 403
    active = client.post(f"/api/oracle/profiles/{profile['id']}/activate").json()
    assert active["status"] == "active" and active["approved_by"]

    # A change is a new draft version; the approval doesn't carry over.
    spec = _spec_for(template)
    spec["mappings"].append({"column": "CurrencyCode", "source": {"kind": "constant", "constant": "USD"}})
    saved = client.put(f"/api/oracle/profiles/{profile['id']}",
                       json={"name": profile["name"], "spec": spec}).json()
    assert (saved["status"], saved["version"], saved["approved_by"]) == ("draft", 2, None)

    compared = client.get(f"/api/oracle/profiles/{profile['id']}/compare", params={"from_version": 1, "to_version": 2}).json()
    assert [(c["column"], c["change"]) for c in compared["changes"]] == [("CurrencyCode", "added")]
    versions = client.get(f"/api/oracle/profiles/{profile['id']}/versions").json()
    assert [v["version"] for v in versions] == [1, 2] and versions[0]["status"] == "active"

    copy = client.post(f"/api/oracle/profiles/{profile['id']}/duplicate").json()
    assert copy["id"] != profile["id"] and (copy["status"], copy["version"]) == ("draft", 1)
    archived = client.post(f"/api/oracle/profiles/{profile['id']}/archive").json()
    assert archived["status"] == "archived"

    history = [e["action"] for e in client.get(f"/api/oracle/profiles/{profile['id']}/history").json()]
    assert history == ["archived", "reverted_to_draft", "activated", "created"]


def test_approval_needs_a_verified_schema_and_every_required_attribute_mapped() -> None:
    draft_target = _register(application="Enterprise Contracts", integration_method="fbdi")
    profile = client.post(f"/api/oracle/templates/{draft_target['id']}/profiles",
                          json={"name": "Draft target", "spec": _spec_for(draft_target)}).json()
    refused = client.post(f"/api/oracle/profiles/{profile['id']}/activate")
    assert refused.status_code == 409 and "draft" in refused.text

    verified = _verified()
    incomplete = client.post(f"/api/oracle/templates/{verified['id']}/profiles", json={"name": "No number", "spec": {
        "mappings": [{"column": "ExtKey", "source": {"kind": "constant", "constant": "K1"}}]}}).json()
    refused = client.post(f"/api/oracle/profiles/{incomplete['id']}/activate")
    assert refused.status_code == 409 and "ContractNumber is required" in refused.text


def test_an_approved_profile_stays_on_its_schema_version() -> None:
    template = _verified()
    profile = client.post(f"/api/oracle/templates/{template['id']}/profiles",
                          json={"name": "Pinned", "spec": _spec_for(template)}).json()
    client.post(f"/api/oracle/profiles/{profile['id']}/activate")
    # A newer version of the same object is registered.
    newer = client.post("/api/oracle/templates",
                        files={"file": ("Contract.csv", _template_bytes() + b"\n#v2", "text/csv")},
                        data={"object_name": template["object_name"]}).json()
    assert newer["version"] == template["version"] + 1
    listed = next(p for p in client.get(f"/api/oracle/templates/{template['id']}/profiles").json() if p["id"] == profile["id"])
    assert listed["template_id"] == template["id"] and listed["status"] == "active"
    assert listed["schema_update_available"] and listed["template_version"] == template["version"]
    assert client.get(f"/api/oracle/templates/{template['id']}").json()["catalog"]["status"] == "stale"


# ---------------------------------------------------------------- 12. tenant isolation


def test_targets_and_profiles_of_one_workspace_are_invisible_to_another() -> None:
    from app.core.auth import workspace_key
    from app.database.session import SessionLocal
    from app.models.oracle_template import OracleMappingProfile, OracleTemplate

    a, b = {"X-Workspace-Token": secrets.token_hex(32)}, {"X-Workspace-Token": secrets.token_hex(32)}
    database = SessionLocal()
    try:
        template = OracleTemplate(object_name=f"Isolated {uuid4().hex[:6]}", version=1, filename="t.csv", file_format="csv",
                                  original_bytes=b"A,B\n", checksum_sha256=uuid4().hex, technical_row=0, columns=[
                                      {"position": 0, "technical_name": "A", "display_label": None, "required": None}],
                                  warnings=[], status="active", owner_workspace=workspace_key(a["X-Workspace-Token"]),
                                  uploaded_by="a", catalog={"application": "Procurement", "integration_method": "fbdi"})
        database.add(template)
        database.flush()
        profile = OracleMappingProfile(template_id=template.id, name="A's", mappings=[], key_columns=[], version=1, history=[],
                                       owner_workspace=template.owner_workspace, updated_by="a", status="draft")
        database.add(profile)
        database.commit()
        template_id, profile_id = template.id, profile.id
    finally:
        database.close()
    settings = get_settings().model_copy(update={"rbac_enforced": True, "public_workspace_access": True})
    with patch("app.core.auth.get_settings", return_value=settings):
        mine = client.get("/api/oracle/catalog", headers=a).json()
        theirs = client.get("/api/oracle/catalog", headers=b).json()
        assert any(o["template"]["id"] == template_id for g in mine["applications"] for o in g["objects"])
        assert not any(o["template"]["id"] == template_id for g in theirs["applications"] for o in g["objects"])
        assert client.get(f"/api/oracle/templates/{template_id}", headers=b).status_code == 404
        assert client.get(f"/api/oracle/profiles/{profile_id}/versions", headers=b).status_code == 404
        # No credentials are ever part of what a workspace sees.
        assert "password" not in json.dumps(mine).lower()


# ---------------------------------------------------------------- 13. preview


def test_preview_shows_source_and_target_values_states_transforms_and_the_parent() -> None:
    document_id = _document()
    header = _verified()
    line = _register(application="Enterprise Contracts", integration_method="fbdi",
                     content=b"HeaderExtKey,LineNumber,Amount\n,,\n#" + uuid4().hex.encode())
    header_profile = client.post(f"/api/oracle/templates/{header['id']}/profiles", json={"name": "Header", "spec": {
        "mappings": [{"column": "ExtKey", "source": _pair("contract.contract_number").model_dump()},
                     {"column": "ContractNumber", "source": _pair("contract.contract_number").model_dump()}]}}).json()
    spec = {
        "row_dataset": "line_items",
        "mappings": [
            {"column": "HeaderExtKey", "source": _pair("contract.contract_number").model_dump()},
            {"column": "LineNumber", "source": {"kind": "field", "dataset_id": "line_items", "field": "contract.line.item_number", "shape": "column"}},
            {"column": "Amount", "transform": "currency", "options": {"decimals": 2},
             "source": {"kind": "field", "dataset_id": "line_items", "field": "contract.line.amount", "shape": "column"}},
        ],
        "parent": {"profile_id": header_profile["id"], "links": [["HeaderExtKey", "ExtKey"]]},
    }
    with _workbook_for_documents(_workbook()):
        preview = client.post(f"/api/documents/{document_id}/oracle/templates/{line['id']}/preview", json={"spec": spec})
    assert preview.status_code == 200, preview.text
    body = preview.json()
    amount = body["headers"].index("Amount")
    assert (body["rows"][0][amount]["source_value"], body["rows"][0][amount]["value"]) == ("USD 1,420,351.70", "1420351.70")
    assert body["transforms"]["Amount"] == "Currency normalization to 2 decimals"
    assert body["column_states"]["LineNumber"] == "Valid"
    assert body["schema_status"] == "draft"
    assert any(i["code"] == "schema_unverified" for i in body["issues"])
    assert body["parent"]["name"] == "Header" and body["parent"]["rows"][0][1]["value"] == "FA300224C0008"


def test_clauses_are_a_mapping_source_by_regulation_family() -> None:
    workbook = _workbook()
    workbook.datasets.append(_dataset(
        "clauses", "Clauses", "repeating",
        [("number", "clause.number", "Clause Number"), ("title", "clause.title", "Title")],
        [["52.204-21", "Basic Safeguarding"], ["252.204-7012", "Safeguarding Covered Defense Information"],
         ["5352.201-9101", "Ombudsman"], ["52.204-21", "Basic Safeguarding"]],
    ))
    records = clause_records(workbook)
    assert [(r.clause_id, r.regulation_family) for r in records] == [
        ("52.204-21", "FAR"), ("252.204-7012", "DFARS"), ("5352.201-9101", "Agency supplement")]
    assert records[0].contract_number == "FA300224C0008" and records[0].title == "Basic Safeguarding"
    document_id = _document()
    with _workbook_for_documents(workbook):
        body = client.get(f"/api/documents/{document_id}/oracle/clauses").json()
    assert body["families"] == {"FAR": 1, "DFARS": 1, "Agency supplement": 1}


# ---------------------------------------------------------------- 14. not configured, discovery


def test_without_a_template_or_oracle_connection_the_workspace_says_so() -> None:
    with patch("app.api.oracle_templates._visible", side_effect=lambda statement, model, actor: statement.where(model.id.is_(None))):
        empty = client.get("/api/oracle/catalog").json()
    assert (empty["configured"], empty["message"]) == (False, "Oracle import template not configured")
    assert empty["discovery"]["configured"] is False
    refused = client.post("/api/oracle/discovery/register", json={"resource": "contracts", "application": "Enterprise Contracts"})
    assert refused.status_code == 409


class _FakeOracle:
    name = "fusion_rest_describe"
    configured = True
    message = "Oracle REST metadata (test)."

    def describe(self, resource: str):
        payload = {"Resources": {resource: {"title": "Contracts", "children": {"ContractLine": {}}, "attributes": [
            {"name": "ContractNumber", "type": "string", "mandatory": True, "maxLength": "120", "title": "Contract Number"},
            {"name": "ContractAmount", "type": "number", "mandatory": False, "precision": 18, "title": "Amount"},
            {"name": "StartDate", "type": "date", "title": "Start Date", "lov": {"childRef": "x"}},
        ]}}}
        found = discovery.parse_describe(payload, resource)
        found.raw = json.dumps(payload).encode()
        return found


def test_discovered_metadata_registers_a_draft_rest_target_with_typed_attributes() -> None:
    with patch("app.api.oracle_templates.discovery.provider", return_value=_FakeOracle()):
        response = client.post("/api/oracle/discovery/register",
                               json={"resource": "contracts", "application": "Enterprise Contracts",
                                     "object_name": f"Contracts REST {uuid4().hex[:6]}"})
    assert response.status_code == 201, response.text
    template = response.json()
    assert template["catalog"]["source"] == "metadata_api" and template["catalog"]["integration_method"] == "rest"
    assert template["catalog"]["status"] == "draft" and template["catalog"]["object_code"] == "contracts"
    columns = {c["technical_name"]: c for c in template["columns"]}
    assert columns["ContractNumber"]["required"] is True and columns["ContractNumber"]["max_length"] == 120
    assert columns["ContractAmount"]["precision"] == 18 and columns["StartDate"]["data_type"] == "date"
    # A REST target is mapped and previewed, never written as an import file.
    document_id = _document()
    with _workbook_for_documents(_workbook()):
        export = client.post(f"/api/documents/{document_id}/oracle/templates/{template['id']}/export",
                             json={"spec": {"mappings": []}})
    assert export.status_code == 409


def test_discovery_is_off_for_the_example_host_and_without_credentials() -> None:
    base = get_settings().model_copy(update={"oracle_base_url": "https://example.fa.oraclecloud.com",
                                             "oracle_metadata_username": None, "oracle_metadata_password": None})
    assert discovery.provider(base.model_copy(update={"oracle_metadata_username": "u", "oracle_metadata_password": "p"})).configured is False
    real = base.model_copy(update={"oracle_base_url": "https://acme.fa.us2.oraclecloud.com"})
    assert discovery.provider(real).configured is False
    found = discovery.provider(real.model_copy(update={"oracle_metadata_username": "u", "oracle_metadata_password": "p"}))
    assert found.configured and "acme.fa.us2.oraclecloud.com" in found.message and "p" not in found.message.split()
    with pytest.raises(discovery.DiscoveryError):
        found.describe("../../etc")
    with pytest.raises(discovery.DiscoveryError):
        discovery.parse_describe({"Resources": {"x": {"attributes": []}}}, "x")


def test_existing_templates_without_a_catalog_read_as_unverified_drafts() -> None:
    class Legacy:
        catalog = None
        status = "active"
        columns: list = []

    assert status_of(Legacy()) == ("draft", "Its Oracle application and integration method aren't set.")

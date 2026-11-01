"""Original Structure: fields and tables in document order, joined to
extraction and review, with disagreements flagged."""

from __future__ import annotations

from uuid import uuid4

import fitz
import pytest
from fastapi.testclient import TestClient

from app.database.session import SessionLocal
from app.main import app
from app.models.document import Document
from app.models.document_metadata_field import DocumentMetadataField
from app.source_structure.original_layout import build_original_layout
from app.source_structure.pdf_structure import extract_page_structure, native_words
from app.source_structure.service import get_or_build_source_structure
from contract_brief_fixture import MOD_COUNT, MOD_HEADERS, contract_brief_pdf, expected_fields

client = TestClient(app)


def _upload(content: bytes, name: str) -> str:
    response = client.post("/api/documents/upload", files={"file": (f"{name}-{uuid4().hex[:6]}.pdf", content, "application/pdf")})
    assert response.status_code == 201, response.text
    document_id = response.json()["document_id"]
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}).status_code == 200
    return document_id


@pytest.fixture(scope="module")
def brief_id() -> str:
    return _upload(contract_brief_pdf(), "brief")


def _layout(document_id: str, metadata: list[DocumentMetadataField] | None = None) -> dict:
    database = SessionLocal()
    try:
        structure = get_or_build_source_structure(database, database.get(Document, document_id))
        return build_original_layout(structure, metadata or [])
    finally:
        database.close()


def _fields(layout: dict) -> list[dict]:
    return [item for section in layout["sections"] for item in section["items"] if item["kind"] == "field"]


def _tables(layout: dict) -> list[dict]:
    return [item for section in layout["sections"] for item in section["items"] if item["kind"] == "table"]


def test_every_printed_field_is_kept_in_document_order(brief_id: str) -> None:
    fields = _fields(_layout(brief_id))
    pairs = [(f["label"], f["document_value"]) for f in fields]
    # Each expected pair is found after the one before it ("Address" repeats
    # under every agency, so order is checked by walking forward).
    position = -1
    for _, label, value in expected_fields():
        position = pairs.index((label.strip(), value), position + 1)
    # Blank form fields are shown, not dropped.
    blank = {f["label"] for f in fields if f["blank"]}
    assert {"9a. Prime Contractor", "9b. Prime Contract No."} <= blank


def test_sub_labels_keep_their_group_and_original_labels(brief_id: str) -> None:
    fields = _fields(_layout(brief_id))
    contract = [f for f in fields if f["group"] == "2. Contract Information"]
    assert [f["label"] for f in contract] == [
        "Contract Number", "Contract Name", "Oracle Project Number", "Date of Award",
        "Security Classification Level", "DPAS Rating",
    ]
    limitations = next(f for f in fields if f["label"] == "If yes, identify the limitations.")
    assert limitations["group"].startswith("23. Does the contract contain ceilings")
    assert limitations["document_value"] == "$40,000,000.00"


def test_sections_follow_the_document_headings(brief_id: str) -> None:
    layout = _layout(brief_id)
    titles = [section["title"] for section in layout["sections"]]
    # Headings are sections, never business fields.
    labels = {f["label"] for f in _fields(layout)} | {f["document_value"] for f in _fields(layout)}
    assert not labels & set(titles)
    assert titles == [
        "A. General Information",
        "B. Contract Clauses and Special Provisions",
        "Contract Summary",
        "C. Summary of Contract Modifications",
    ]


def test_tables_keep_printed_headers_and_merge_continued_pages_only_on_evidence(brief_id: str) -> None:
    summary, modifications, accruals = _tables(_layout(brief_id))
    assert summary["headers"][:3] == ["Period Name", "Exercised?", "Start"]
    assert modifications["headers"] == MOD_HEADERS
    assert modifications["pages"] == [4, 5] and modifications["continued"]
    # Every row once, in printed order, as rows (not one row per cell).
    assert [int(row["cells"][0]) for row in modifications["rows"]] == list(range(1, MOD_COUNT + 1))
    assert len(modifications["rows"][0]["cells"]) == len(MOD_HEADERS)
    # Page 6 repeats the header but its rows restart: a different table.
    assert accruals["pages"] == [6] and not accruals["continued"]
    assert [row["cells"][0] for row in accruals["rows"]] == ["1", "2"]


def _metadata(document_id: str, key: str, label: str, value: str, page: int, *, status: str = "pending", reviewed: str | None = None) -> DocumentMetadataField:
    return DocumentMetadataField(
        id=abs(hash((key, value, page))) % 10**6, document_id=document_id, field_group="g", field_key=key, label=label,
        value=reviewed if reviewed is not None else value, original_value=value, confidence=0.9,
        extraction_method="target", review_status=status,
        evidence_json={"page_number": page, "source_text": f"{label} {value}", "machine_value": value},
        extraction_source="target",
    )


def test_label_value_disagreements_are_flagged_and_reviewed_values_kept(brief_id: str) -> None:
    rows = [
        _metadata(brief_id, "cage_code", "CAGE Code", "7QRT2", 1, status="accepted"),
        _metadata(brief_id, "contract_no", "Contract Number", "", 1),  # the 25% "—" row
        _metadata(brief_id, "contract_number", "Contract Number", "C05550001", 1),
        _metadata(brief_id, "business_size", "Business Size", "Small", 1),  # wrong value for the label
        _metadata(brief_id, "role", "Company Role", "Prime", 1, status="edited", reviewed="Prime Contractor"),
        _metadata(brief_id, "contract_type", "Contract Type", "DIRECT", 1),
        _metadata(brief_id, "contract_type", "Contract Type", "INDIRECT", 4),
    ]
    by_label = {f["label"]: f for f in _fields(_layout(brief_id, rows)) if f["page"] in (1, 4)}

    assert by_label["Cage"]["review_status"] == "Approved" and not by_label["Cage"]["flags"]
    number = next(f for f in _fields(_layout(brief_id, rows)) if f["label"] == "Contract Number" and f["page"] == 1)
    assert number["extracted_value"] == "C05550001"
    assert [f["code"] for f in number["flags"]] == ["missed_value"]
    assert number["review_status"] == "Needs review"
    # Still for review, but the value was read: it isn't "missing".
    assert number["quality"] == "confirmed"
    assert "label_value_mismatch" in [f["code"] for f in by_label["Business Size"]["flags"]]
    assert by_label["Business Size"]["document_value"] == "HUBZone"  # never replaced
    role = by_label["8. Company Role"]
    assert (role["extracted_value"], role["effective_value"], role["review_status"]) == ("Prime", "Prime Contractor", "Corrected")
    assert "conflicting_values" in [f["code"] for f in by_label["4. Contract Type"]["flags"]]


def test_table_cells_that_dont_fit_their_column_are_flagged() -> None:
    from app.source_structure.models import TableCandidate, TableCell, StructuredSourceDocument, StructureStats

    rows = [["1", "$1,000.00"], ["2", "$2,000.00"], ["3", "#########"], ["4", "$4,000.00"], ["5", "$5,000.00"]]
    table = TableCandidate(
        candidate_id="t1", region_id="t1", page=1, bbox=(0, 100, 100, 200), detection_method="pdf_ruling_lines",
        headers=["No", "Amount"], header_cells=[TableCell(row_index=-1, column_index=i, text=h) for i, h in enumerate(["No", "Amount"])],
        rows=[[TableCell(row_index=r, column_index=c, text=t) for c, t in enumerate(row)] for r, row in enumerate(rows)],
    )
    structure = StructuredSourceDocument(document_id="d", source_type="pdf", extractor_version=0, regions=[], field_candidates=[], table_candidates=[table], stats=StructureStats(pages=1))
    [section] = build_original_layout(structure, [])["sections"]
    flagged = [row for row in section["items"][0]["rows"] if row.get("flags")]
    assert [row["cells"][1] for row in flagged] == ["#########"]
    assert flagged[0]["review_status"] == "Needs review"


def test_box_borders_are_not_read_as_fill_in_underlines() -> None:
    doc = fitz.open()
    page = doc.new_page()
    for index, (label, value) in enumerate([("7. ISSUED BY", "N40192"), ("8. ADDRESS", "Code 1"), ("9. NAME", "Paul"), ("10. PHONE", "210")]):
        top = 80 + index * 40
        page.draw_rect(fitz.Rect(60, top, 540, top + 30), width=0.6)  # a boxed form cell
        page.insert_text((64, top + 12), label, fontsize=8)
        page.insert_text((300, top + 26), value, fontsize=9)
    result = extract_page_structure(page_number=1, words=native_words(page), extraction_method="native", fitz_page=page)
    assert not [f for f in result.fields if f.structural_relation == "ruled_form_row"]


def _hairline(page: fitz.Page, x0: float, y0: float, x1: float, y1: float) -> None:
    # Spreadsheet exports draw each rule as a thin filled rectangle.
    page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=None, fill=(0, 0, 0))


def test_rectangle_drawn_grid_below_a_label_block() -> None:
    """aa.pdf page 4: underlined label/value rows over a grid whose rules are
    all thin filled rectangles; long values run past their underline."""

    doc = fitz.open()
    page = doc.new_page()
    block = [("Contractor Name", "ORBITAL FIELD SERVICES, LLC"), ("Contract Name", "OFS JV"), ("Contract Number", "C05550001"), ("Contract Type", "DIRECT")]
    for index, (label, value) in enumerate(block):
        y = 90 + index * 12
        page.insert_text((60, y), label, fontsize=7)
        page.insert_text((160, y), value, fontsize=7)
        _hairline(page, 158, y + 1.5, 230, y + 2)  # shorter than the long name
    columns = [60, 110, 200, 280, 360]
    rows = [["Mod", "Description", "Date", "Funding"]] + [[str(i), "AWARD" if i == 1 else "", f"07/{i:02d}/2025", f"${i},000.00"] for i in range(1, 7)]
    top = 160
    for r, cells in enumerate(rows):
        y = top + r * 12
        for c, text in enumerate(cells):
            page.insert_text((columns[c] + 2, y + 9), text, fontsize=7)
    for r in range(len(rows) + 1):
        _hairline(page, columns[0], top + r * 12, columns[-1], top + r * 12 + 0.5)
    for x in columns:
        _hairline(page, x, top, x + 0.5, top + len(rows) * 12)

    result = extract_page_structure(page_number=1, words=native_words(page), extraction_method="native", fitz_page=page)
    fields = {f.label_text: f.raw_value for f in result.fields if f.structural_relation == "ruled_form_row"}
    assert fields["Contractor Name"] == "ORBITAL FIELD SERVICES, LLC"
    assert fields["Contract Type"] == "DIRECT"
    [table] = result.tables
    assert table.headers == ["Mod", "Description", "Date", "Funding"]
    assert len(table.rows) == 6
    assert not [r for r in result.regions if r.region_type == "HEADING"]


def _export_values(brief_id: str):
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get(f"/api/documents/{brief_id}/original-structure/export.xlsx")
    assert response.status_code == 200
    assert "original file.xlsx" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    sheet = book.active
    rows = list(sheet.iter_rows())
    values = [[c.value for c in row] for row in rows]
    return book, sheet, rows, values


def test_original_document_structure_export_is_one_editable_sheet_in_document_order(brief_id: str) -> None:
    from contract_brief_fixture import SUMMARY_HEADERS

    book, sheet, rows, values = _export_values(brief_id)
    # One worksheet; no merged cell anywhere, so every value can be edited,
    # sorted and filtered.
    assert book.sheetnames == ["Original File"]
    assert not sheet.merged_cells.ranges
    # Every value is text exactly as printed: no dates or numbers made by Excel.
    assert {c.data_type for row in rows for c in row if c.value is not None} == {"s"}

    # A section opens with its printed heading; its fields sit under the
    # field columns.
    titles = [r[0] for r in values if r[0] and str(r[0]).startswith(("A. ", "B. ", "C. "))]
    assert any(t.startswith("C. Summary of Contract Modifications") for t in titles)
    header = ["Item #", "Field Name", "Extracted Value"]
    assert values.index(next(r for r in values if r[:3] == header)) > 0

    # Item # | Field Name | Extracted Value: an item with sub-fields opens
    # with its number and name on a row of their own; each sub-field follows,
    # its number cell blank. Nothing is merged into one cell.
    first = next(i for i, r in enumerate(values) if r[:2] == ["10.", "Acquisition Agency"])
    assert values[first][2] is None
    assert values[first + 1][:3] == [None, "Agency", "DOD Space Force"]
    assert values[first + 2][:3] == [None, "Address", None]
    assert sum(1 for r in values if r[0] == "10.") == 1
    assert sum(1 for r in values if r[1] == "Acquisition Agency") == 1
    # An item that is one field is one row.
    briefed = next(r for r in values if r[0] == "3.")
    assert briefed[1:3] == ["Briefed Through", "MOD#12 - 07/15/2026"]

    # Native tables keep their printed headers, columns and row order, with
    # no columns added.
    summary = next(i for i, r in enumerate(values) if r[0] == SUMMARY_HEADERS[0])
    assert [v for v in values[summary] if v is not None] == SUMMARY_HEADERS
    mods = [i for i, r in enumerate(values) if r[0] == MOD_HEADERS[0]]
    assert len(mods) == 2 and all([v for v in values[i] if v is not None] == MOD_HEADERS for i in mods)
    numbers = [r[0] for r in values[mods[0] + 1:] if r[0] and str(r[0]).isdigit()][:MOD_COUNT]
    assert numbers == [str(n) for n in range(1, MOD_COUNT + 1)]

    # The summary table names its source page and carries its contract's
    # metadata once above it; the modification tables have their own block
    # printed above them, so nothing is added there.
    assert values[summary - 2][:2] == ["Table", "Source: page 3"]
    assert values[summary - 1][:2] == ["   ·   ".join([
        "Legal Entity: ORBITAL FIELD SERVICES, LLC", "Contract Number: C05550001",
        "Oracle Project Number: 90TTSC026C0002", "Contract Type: DIRECT"]), None]
    contexts = [r for r in values if r[0] and str(r[0]).startswith("Legal Entity: ")]
    assert len(contexts) == 1
    assert all(not str(values[i - 1][0] or "").startswith("Legal Entity:") for i in mods)


def test_the_export_loses_no_record(brief_id: str) -> None:
    from app.source_structure.original_export import original_rows

    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    rows = original_rows(layout)
    # One row per field (placed or not) and per table row, in order.
    assert sum(1 for r in rows if r.kind in ("field", "subfield")) == len(_fields(layout)) + len(layout["unplaced"])
    assert [r.values for r in rows if r.kind == "table_row"] == [
        row.get("reviewed_cells") or row["cells"] for t in _tables(layout) for row in t["rows"]
    ]
    _, _, _, values = _export_values(brief_id)
    assert len(values) == len(rows)


def test_numbered_items_are_decided_once_in_the_layout(brief_id: str) -> None:
    from app.source_structure.original_layout import item_parts

    assert item_parts(None, "3. Briefed Through") == ("3.", "Briefed Through", "Briefed Through", False)
    assert item_parts("1. Contractor Information", "Cage") == ("1.", "Contractor Information", "Cage", True)
    assert item_parts(None, "Contract Type") == ("", "", "Contract Type", False)

    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    fields = _fields(layout)
    agency = [f for f in fields if (f.get("item") or {}).get("name") == "Acquisition Agency"]
    assert [f["field_name"] for f in agency] == ["Agency", "Address", "Point of Contact", "Email", "Phone"]
    assert [f["item_start"] for f in agency] == [True, False, False, False, False]
    # Children point at their parent by a stable id; the item at its section.
    assert {f["parent_id"] for f in agency} == {agency[0]["item"]["id"]}
    again = client.get(f"/api/documents/{brief_id}/original-structure").json()
    assert [f["parent_id"] for f in _fields(again)] == [f["parent_id"] for f in fields]
    for section in layout["sections"]:
        for item in section["items"]:
            if item["kind"] == "table":
                assert item["parent_id"] == section["id"]
                assert all(row["parent_id"] == item["record_key"] for row in item["rows"])


def test_table_metadata_belongs_to_the_contract_printed_above_it(brief_id: str) -> None:
    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    tables = _tables(layout)
    context = [(t["context_source"], {e["label"]: e["value"] for e in t["context"]}) for t in tables]
    assert context[0][0] == "inherited" and context[0][1]["Contract Number"] == "C05550001"
    assert context[1][0] == "printed" and context[1][1]["Contract Number"] == "C05550001"
    # The accrual contract's table: its own number; its blank contractor
    # stays blank rather than borrowing the first contract's.
    assert context[2][0] == "printed" and context[2][1]["Contract Number"] == "C05550001.00.9999"
    assert context[2][1]["Contractor Name"] == ""
    # The concept, and the label exactly as printed.
    assert layout["document_context"][0] == {
        "label": "Contractor Name", "printed_label": "Legal Entity", "value": "ORBITAL FIELD SERVICES, LLC",
        "record_key": layout["document_context"][0]["record_key"], "page": 1,
    }


def test_data_quality_is_one_state_per_value_and_blank_is_not_missing(brief_id: str) -> None:
    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    fields = _fields(layout)
    assert all(f["quality"] in layout["quality_labels"] for f in fields)
    blank = next(f for f in fields if f["blank"])
    assert blank["quality"] == "blank" and blank["quality_label"] == "Blank in source"
    printed = next(f for f in fields if f["label"] == "Cage")
    assert printed["quality"] == "confirmed"
    assert sum(layout["stats"]["quality"].values()) == len(fields) + len(layout["unplaced"])


def test_quality_states_follow_the_source_and_extraction() -> None:
    from app.source_structure.models import FieldCandidate
    from app.source_structure.original_layout import _field_quality

    def candidate(**extra):
        return FieldCandidate.model_construct(**{"extraction_method": "native", "ocr_confidence": None, "ocr_contested": False, **extra})

    def payload(printed, extracted, flags=(), review="Not reviewed"):
        return {"document_value": printed, "extracted_value": extracted, "review_status": review,
                "flags": [{"code": code} for code in flags]}

    assert _field_quality(payload("C05550001", "C05550001"), candidate()) == "confirmed"
    assert _field_quality(payload("", ""), candidate()) == "blank"
    assert _field_quality(payload("#####", "#####"), candidate()) == "unreadable"
    assert _field_quality(payload("A\ufffdB", "A\ufffdB"), candidate()) == "unreadable"
    # Read from the layout with no extraction row: the printed value stands.
    assert _field_quality(payload("FAR", ""), candidate()) == "confirmed"
    assert _field_quality(payload("FAR", "", ["missed_value"]), candidate()) == "missing_extraction"
    assert _field_quality(payload("FAR", "FAR", ["missed_value"]), candidate()) == "confirmed"
    assert _field_quality(payload("FAR", "DFARS", ["label_value_mismatch"]), candidate()) == "conflicting"
    assert _field_quality(payload("FAR", "FAR", ["conflicting_values"]), candidate()) == "conflicting"
    ocr = candidate(extraction_method="ocr", ocr_confidence=0.62)
    assert _field_quality(payload("FAR", "FAR"), ocr) == "unverified_ocr"
    # A reviewer's approval verifies an OCR read; a confident read needs none.
    assert _field_quality(payload("FAR", "FAR", review="Approved"), ocr) == "confirmed"
    sure = candidate(extraction_method="ocr", ocr_confidence=0.97)
    assert _field_quality(payload("FAR", "FAR"), sure) == "confirmed"
    contested = candidate(extraction_method="ocr", ocr_confidence=0.97, ocr_contested=True)
    assert _field_quality(payload("FAR", "FAR"), contested) == "unverified_ocr"
    # Extracted with no printed position: never claimed as a confirmed source value.
    # Extracted with no printed position: neither confirmed nor called wrong.
    assert _field_quality(payload(None, "C05550001"), None) == "extracted"
    assert _field_quality(payload(None, ""), None) == "missing_extraction"


def test_values_that_are_not_confirmed_carry_a_note_and_are_never_changed() -> None:
    from app.source_structure.original_export import original_rows

    def field(label, value, quality, **extra):
        return {"kind": "field", "record_key": label, "label": label, "field_name": label, "item": None,
                "document_value": value, "extracted_value": value, "effective_value": value,
                "quality": quality, "other_metadata": [], **extra}

    layout = {"sections": [{"title": "A. Header", "level": 1, "items": [
        field("Cage", "7QRT2", "confirmed"),
        field("Fax", "", "blank"),
        field("Phone", "#####", "unreadable"),
        field("Email", "a@b.mil", "unverified_ocr", ocr_confidence=0.61),
        {"kind": "table", "record_key": "table:2:0", "title": "Rates", "headers": ["Labor", "Rate"], "pages": [2],
         "context": [], "context_source": None,
         "rows": [{"record_key": "table:2:0:row:0", "cells": ["Analyst", "#####"],
                   "cell_issues": [{"column": 1, "quality": "unreadable", "message": "Unreadable source value"}]}]},
    ]}], "unplaced": []}
    rows = original_rows(layout)
    by_label = {row.values[1]: row for row in rows if row.kind in ("field", "subfield")}
    assert by_label["Cage"].notes == {} and by_label["Fax"].notes == {} and by_label["Fax"].values[2] == ""
    assert by_label["Phone"].values[2] == "#####" and by_label["Phone"].notes[2] == "Unreadable source value"
    assert by_label["Email"].notes[2] == "Unverified OCR result (OCR confidence 0.61)"
    table_row = next(row for row in rows if row.kind == "table_row")
    assert table_row.values == ["Analyst", "#####"] and table_row.notes == {1: "Unreadable source value"}
    # The table title differs from its section's, so it is kept as printed.
    assert next(row for row in rows if row.kind == "table_title").values == ["Rates", "Source: page 2"]


def _row(cells, **extra):
    return {"record_key": f"t:row:{cells[0]}", "cells": cells, **extra}


def test_a_date_range_is_split_only_when_both_dates_are_certain() -> None:
    from app.source_structure.original_layout import composite_parts

    assert composite_parts("01/01/2025 - 12/31/2025") == [
        {"label": "Start Date", "value": "01/01/2025"}, {"label": "End Date", "value": "12/31/2025"}]
    assert [p["value"] for p in composite_parts("July 1, 2025 through June 30, 2026")] == ["July 1, 2025", "June 30, 2026"]
    assert [p["value"] for p in composite_parts("2025-07-01 to 2026-06-30")] == ["2025-07-01", "2026-06-30"]
    # One date, an open end, or other text: left whole, nothing invented.
    for whole in ["01/01/2025", "01/01/2025 - TBD", "Base period - 12 months", "C05550001 - C05550002", "", None]:
        assert composite_parts(whole) == []


def test_a_split_value_keeps_its_printed_text_and_writes_each_part_once() -> None:
    from app.source_structure.original_export import original_rows

    period = {"kind": "field", "record_key": "pop", "label": "Period of Performance", "field_name": "Period of Performance",
              "item": None, "document_value": "01/01/2025 - 12/31/2025", "extracted_value": "01/01/2025 - 12/31/2025",
              "effective_value": "01/01/2025 - 12/31/2025", "quality": "confirmed", "other_metadata": [],
              "parts": [{"label": "Start Date", "value": "01/01/2025"}, {"label": "End Date", "value": "12/31/2025"}]}
    rows = original_rows({"sections": [{"title": "B. Contract", "level": 1, "items": [period]}], "unplaced": []})
    parent = next(r for r in rows if r.kind == "field")
    assert parent.values == ["", "Period of Performance", ""]
    assert parent.notes == {2: "Printed as 01/01/2025 - 12/31/2025; split into Start Date and End Date"}
    assert [r.values for r in rows if r.kind == "part"] == [
        ["", "Start Date", "01/01/2025"], ["", "End Date", "12/31/2025"]]


def test_a_column_naming_one_contract_for_every_row_is_shown_once_above_the_table() -> None:
    from app.source_structure.original_export import original_rows
    from app.source_structure.original_layout import _shared_columns

    headers = ["Item", "Contract Number", "Period", "Start Date", "End Date", "Amount"]
    rows = [_row(["1", "FA30-001", "Base", "01/01/2025", "12/31/2025", "$100,000"]),
            _row(["1", "FA30-001", "Option 1", "01/01/2026", "", "$110,000"]),
            _row(["1", "FA30-001", "Option 2", "01/01/2027", "12/31/2027", "$120,000"])]
    shared = _shared_columns(headers, rows, printed=True)
    assert shared == [{"column": 0, "label": "Item", "value": "1"}, {"column": 1, "label": "Contract Number", "value": "FA30-001"}]
    # A value that changes between rows stays in the rows; so does a
    # constant column that names no group, and a one-row table.
    changing = rows[:2] + [_row(["2", "FA30-002", "Base", "", "", "$1"])]
    assert _shared_columns(headers, changing, printed=True) == []
    assert _shared_columns(["Period", "Currency"], [_row(["Base", "USD"]), _row(["Option", "USD"])], printed=True) == []
    assert _shared_columns(headers, rows[:1], printed=True) == []
    assert _shared_columns(headers, rows, printed=False) == []

    table = {"kind": "table", "record_key": "t", "title": "Contract Periods", "headers": headers, "pages": [2],
             "context": [], "context_source": None, "shared_columns": shared,
             "rows": [*rows[:1], {**rows[1], "cell_issues": [{"column": 5, "quality": "unreadable", "message": "Unreadable source value"}]}, rows[2]]}
    out = original_rows({"sections": [{"title": "C. Periods", "level": 1, "items": [table]}], "unplaced": []})
    assert next(r for r in out if r.kind == "context").values == ["Item: 1   ·   Contract Number: FA30-001"]
    assert next(r for r in out if r.kind == "table_header").values == ["Period", "Start Date", "End Date", "Amount"]
    body = [r for r in out if r.kind == "table_row"]
    # Every record kept, in order, its blank cell blank; notes follow their column.
    assert [r.values for r in body] == [
        ["Base", "01/01/2025", "12/31/2025", "$100,000"],
        ["Option 1", "01/01/2026", "", "$110,000"],
        ["Option 2", "01/01/2027", "12/31/2027", "$120,000"]]
    assert body[1].notes == {3: "Unreadable source value"}
    # The data itself is unchanged: the layout's rows still hold every cell.
    assert table["rows"][0]["cells"][1] == "FA30-001"

    # An inherited block naming the same contract doesn't repeat it.
    from app.source_structure.original_layout import _reconcile_context

    inherited = {**table, "context_source": "inherited", "context": [
        {"label": "Contractor Name", "printed_label": "Contractor Name", "value": "NOVA SPACE SOLUTIONS, LLC", "record_key": "a", "page": 1},
        {"label": "Contract Number", "printed_label": "Contract Number", "value": "FA30-001", "record_key": "b", "page": 1}]}
    _reconcile_context(inherited)
    assert [e["label"] for e in inherited["context"]] == ["Contractor Name"]
    out = original_rows({"sections": [{"title": "C. Periods", "level": 1, "items": [inherited]}], "unplaced": []})
    assert next(r for r in out if r.kind == "context").values == [
        "Contractor Name: NOVA SPACE SOLUTIONS, LLC   ·   Item: 1   ·   Contract Number: FA30-001"]


def test_metadata_is_never_inherited_across_a_contract_boundary() -> None:
    from app.source_structure.original_layout import _reconcile_context

    def table(headers, rows, shared):
        return {"headers": headers, "rows": [_row(r) for r in rows], "shared_columns": shared, "context_source": "inherited",
                "context": [{"label": "Contractor Name", "printed_label": "Vendor", "value": "ALPHA LLC", "record_key": "a", "page": 1},
                            {"label": "Contract Number", "printed_label": "Contract No.", "value": "A-100", "record_key": "b", "page": 1}]}

    # Its rows name another contract: nothing of contract A is carried over.
    other = table(["Contract No.", "Amount"], [["C-300", "$1"], ["C-300", "$2"]], [{"column": 0, "label": "Contract No.", "value": "C-300"}])
    _reconcile_context(other)
    assert other["context"] == [] and other["context_source"] is None
    # Rows of several contracts: each row keeps its own; nothing inherited.
    mixed = table(["Contract Number", "Amount"], [["A-100", "$1"], ["B-200", "$2"]], [])
    _reconcile_context(mixed)
    assert mixed["context"] == [] and mixed["context_source"] is None
    # The same contract: kept, said once.
    same = table(["Contract Number", "Amount"], [["A-100", "$1"], ["A-100", "$2"]], [{"column": 0, "label": "Contract Number", "value": "A-100"}])
    _reconcile_context(same)
    assert [e["label"] for e in same["context"]] == ["Contractor Name"]


def _two_contracts_pdf() -> bytes:
    """Contract A (vendor and number), contract B (an agreement number
    only), then a table naming contract C in its rows and one mixing A and B."""

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 60

    def heading(text):
        nonlocal y
        y += 26
        page.insert_text((72, y), text, fontsize=13, fontname="hebo")
        y += 8

    def line(label, value):
        nonlocal y
        y += 18
        page.insert_text((90, y), label, fontsize=10, fontname="hebo")
        page.insert_text((260, y), value, fontsize=10, fontname="helv")

    def grid(headers, rows):
        nonlocal y
        top, height, xs = y + 14, 20, [72 + i * (468 // len(headers)) for i in range(len(headers))] + [540]
        for r in range(len(rows) + 2):
            page.draw_line((xs[0], top + r * height), (xs[-1], top + r * height), width=0.6)
        for x in xs:
            page.draw_line((x, top), (x, top + (len(rows) + 1) * height), width=0.6)
        for c, text in enumerate(headers):
            page.insert_text((xs[c] + 4, top + 14), text, fontsize=9, fontname="hebo")
        for r, row in enumerate(rows, start=1):
            for c, text in enumerate(row):
                page.insert_text((xs[c] + 4, top + r * height + 14), text, fontsize=9, fontname="helv")
        y = top + (len(rows) + 1) * height + 6

    heading("Contract A Details")
    line("Vendor Name:", "ALPHA LLC")
    line("Contract No.:", "A-100")
    heading("Contract A Periods")
    grid(["Period", "Start Date", "Amount"], [["Base", "01/01/2025", "$1,000"], ["Option 1", "01/01/2026", "$2,000"]])
    heading("Contract B Details")
    line("Agreement Number:", "B-200")
    heading("Contract B Periods")
    grid(["Period", "Start Date", "Amount"], [["Base", "01/01/2027", "$3,000"], ["Option 1", "01/01/2028", "$4,000"]])
    heading("Accruals")
    grid(["Contract No.", "Period", "Amount"], [["C-300", "Base", "$5,000"], ["C-300", "Option 1", "$6,000"]])
    heading("Funding by Contract")
    grid(["Contract Number", "Period", "Amount"], [["A-100", "Base", "$7,000"], ["B-200", "Base", "$8,000"]])
    return doc.tobytes()


def test_each_table_shows_only_its_own_contract_in_a_multi_contract_document() -> None:
    layout = client.get(f"/api/documents/{_upload(_two_contracts_pdf(), 'two-contracts')}/original-structure").json()
    tables = {t["title"]: t for t in _tables(layout)}
    shown = {title: [(e["printed_label"], e["value"]) for e in t["context"]] if t["context_source"] == "inherited" else []
             for title, t in tables.items()}
    # Recognised by alias ("Vendor Name", "Contract No.", "Agreement
    # Number") and shown under the printed label.
    assert shown["Contract A Periods"] == [("Vendor Name", "ALPHA LLC"), ("Contract No.", "A-100")]
    # Contract B printed only its number: A's vendor is not carried into it.
    assert shown["Contract B Periods"] == [("Agreement Number", "B-200")]
    # Rows naming contract C: no inherited B; C once above the table.
    assert shown["Accruals"] == []
    assert [(e["label"], e["value"]) for e in tables["Accruals"]["shared_columns"]] == [("Contract No.", "C-300")]
    # Rows of A and B: each row keeps its own contract number.
    assert shown["Funding by Contract"] == [] and tables["Funding by Contract"]["shared_columns"] == []
    assert [r["cells"][0] for r in tables["Funding by Contract"]["rows"]] == ["A-100", "B-200"]


def test_metadata_aliases_are_configurable_and_printed_labels_are_kept(tmp_path, monkeypatch) -> None:
    import json

    from app.core.config import get_settings
    from app.source_structure import original_layout

    assert original_layout.metadata_concept("Legal Entity:") == "Contractor Name"
    assert original_layout.metadata_concept("Contract No.") == "Contract Number"
    assert original_layout.metadata_concept("Lot") is None
    extra = tmp_path / "aliases.json"
    extra.write_text(json.dumps({"Contract Number": ["Lot"], "Task Order": ["task order no"]}), encoding="utf-8")
    monkeypatch.setattr(get_settings(), "metadata_aliases_path", str(extra))
    original_layout.metadata_aliases.cache_clear()
    try:
        assert original_layout.metadata_concept("Lot") == "Contract Number"
        assert original_layout.metadata_concept("Task Order No.") == "Task Order"
        assert original_layout.metadata_concept("Legal Entity") == "Contractor Name"  # built-ins kept
    finally:
        monkeypatch.undo()
        original_layout.metadata_aliases.cache_clear()


def test_source_location_is_kept_apart_from_data_quality(brief_id: str) -> None:
    from app.source_structure.original_export import original_rows

    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    assert {f["source_location"] for f in _fields(layout)} == {"exact"}
    assert all(r["source_location"] == "exact" for t in _tables(layout) for r in t["rows"])
    # An extracted value with only a page: its own quality, and its location said separately.
    unlocated = {"kind": "field", "record_key": "metadata:1", "label": "Effective Date", "field_name": "Effective Date", "item": None,
                 "document_value": None, "extracted_value": "08/13/2025", "effective_value": "08/13/2025", "quality": "extracted",
                 "source_location": "page", "page": 6, "other_metadata": []}
    row = next(r for r in original_rows({"sections": [], "unplaced": [unlocated]}) if r.kind == "field")
    assert row.values[2] == "08/13/2025"
    assert row.notes == {2: "Source: Source page only (no exact position recorded) (page 6)"}
    assert layout["quality_labels"]["extracted"] == "Extracted value"
    assert "unlocated" not in layout["quality_labels"]


def test_technical_names_read_as_labels_and_printed_wording_is_kept() -> None:
    from app.source_structure.original_layout import display_label

    assert display_label("contract_no") == "Contract Number"
    assert display_label("award_dt") == "Award Date"
    assert display_label("vendor_legal_name") == "Vendor Legal Name"
    assert display_label("periodStart") == "Period Start"
    for printed in ["Contract No.", "CAGE Code", "Period of Performance", "3. Briefed Through", "Cage"]:
        assert display_label(printed) == printed


def test_empty_reader_artifacts_are_left_out_but_real_blanks_and_missing_values_stay() -> None:
    from app.source_structure.original_layout import QUALITY_LABELS, _is_artifact

    def payload(label, value="", **extra):
        return {"label": label, "document_value": value, "effective_value": value, "metadata": None, "flags": [], **extra}

    assert _is_artifact(payload(":")) and _is_artifact(payload("3"))
    assert not _is_artifact(payload("Fax"))  # a printed blank
    assert not _is_artifact(payload("4", "07/01/2025"))  # it has a value
    assert not _is_artifact(payload("", metadata={"id": 1}))  # an extraction row
    assert not _is_artifact(payload("", flags=[{"code": "missed_value"}]))
    assert QUALITY_LABELS["missing_extraction"] == "Missing — Review Required"
    assert QUALITY_LABELS["blank"] == "Blank in source"


def test_a_letterhead_repeated_on_every_page_is_one_section_and_its_table_carries_on() -> None:
    from pathlib import Path

    invoice = Path(__file__).parent / "fixtures" / "invoices" / "f3_multipage_parts.pdf"
    layout = client.get(f"/api/documents/{_upload(invoice.read_bytes(), 'f3')}/original-structure").json()
    # The supplier's name printed atop each page is a running header, not
    # three sections of the same name.
    assert [s["title"] for s in layout["sections"]] == ["Summit Industrial Parts, Inc."]
    items = layout["sections"][0]["items"]
    tables = [i for i in items if i["kind"] == "table"]
    # Pages 1 and 2 are linked by the structure (header repeated, table runs
    # to the page foot) and its line numbers carry on at their step of 10:
    # one table. Page 3 has no such link and stays its own table.
    assert [(t["pages"], t["row_count"]) for t in tables] == [([1, 2], 43), ([3], 7)]
    lines = [r["cells"][0] for r in tables[0]["rows"]]
    assert lines == [str(n) for n in range(10, 440, 10)]
    # The invoice number repeated in each page header is kept, as a record,
    # marked as a repeat.
    repeats = [i for i in items if i.get("repeated_header")]
    assert [(i["label"], i["page"]) for i in repeats] == [("Invoice No.", 2), ("Invoice No.", 3)]


def test_form_instructions_are_removed_from_field_names_and_meaningful_names_kept() -> None:
    from app.source_structure.original_layout import form_label

    # (item number, field name, is a tick box)
    assert form_label("2 Enter the Contract Number") == ("2", "Contract Number", False)
    assert form_label("3 Enter the Solicitation Number") == ("3", "Solicitation Number", False)
    assert form_label("Enter the Contractor Name") == ("", "Contractor Name", False)
    assert form_label("1 Enter the Rating for this solicitation, offer, and award") == ("1", "Rating", False)
    assert form_label("5 Enter the Date Issued (enter 2 digit month, 2 digit day and 4 digit year)") == ("5", "Date Issued", False)
    assert form_label('4 Click here to select, "Negotiated (RFP)" for the Type of Solicitation') == ("4", "Negotiated (RFP)", True)
    assert form_label("15b Enter The Offeror's Telephone Number") == ("15B", "Offeror's Telephone Number", False)
    assert form_label("20 (To be completed by the Government) Enter the accepted dollar amount") == ("20", "Accepted Dollar Amount", False)
    assert form_label("14 Enter the first Amendment Number (if applicable) (the offeror acknowledges receipt of "
                      "amendments to the SOLICITATION for offerors and related documents)") == ("14", "First Amendment Number", False)
    # Not an instruction: kept word for word.
    for name in ["Code of Federal Regulations", "Contract Number", "Selected Items", "Entered Date"]:
        assert form_label(name) == ("", name, False)


def _filled_form_pdf() -> bytes:
    """A one-page fillable form in the shape of an SF33 cover: numbered
    boxes whose descriptions are instructions, a ticked choice, a box with
    no number, and printed text that repeats one box's value."""

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 60), "SOLICITATION, OFFER AND AWARD", fontsize=14, fontname="hebo")

    def box(label, value, rect, kind=fitz.PDF_WIDGET_TYPE_TEXT):
        widget = fitz.Widget()
        widget.field_type = kind
        widget.field_name = f"f{len(list(page.widgets() or []))}"
        widget.field_label = label
        widget.rect = fitz.Rect(*rect)
        widget.field_value = value
        page.add_widget(widget)

    box("1 Enter the Rating for this solicitation, offer, and award", "DO-C3", (400, 80, 500, 94))
    box("Enter the page number for this page", "1", (510, 80, 540, 94))
    box("2 Enter the Contract Number", "FA300224C0008", (72, 110, 220, 124))
    box("3 Enter the Solicitation Number", "FA300223R00090009", (230, 110, 380, 124))
    box('4 Click here to select, "Negotiated (RFP)" for the Type of Solicitation', True, (400, 110, 410, 120), fitz.PDF_WIDGET_TYPE_CHECKBOX)
    box("5 Enter the Date Issued (enter 2 digit month, 2 digit day and 4 digit year)", "21 Apr 2023", (430, 110, 540, 124))
    # The same number the box holds, printed as text inside it.
    page.insert_text((80, 120), "FA300224C0008", fontsize=8)
    return doc.tobytes()


def test_a_filled_form_reads_as_item_number_field_name_and_value() -> None:
    from app.source_structure.original_export import original_rows

    layout = client.get(f"/api/documents/{_upload(_filled_form_pdf(), 'sf33')}/original-structure").json()
    fields = [f for f in _fields(layout) if f["found_by"] == "form widget"]
    by_name = {f["field_name"]: f for f in fields}
    # Every filled box is kept, under the form's own field name.
    assert set(by_name) >= {"Rating", "Page Number", "Contract Number", "Solicitation Number", "Negotiated (RFP)", "Date Issued"}
    assert by_name["Rating"]["effective_value"] == "DO-C3"
    # Item numbers come from the form, in its order: 1, 2, 3, 4, 5.
    numbers = [f["item"]["number"] for f in fields if f.get("item") and f["item_start"]]
    assert numbers == ["1", "2", "3", "4", "5"]
    # The original wording stays on the record, for traceability.
    assert by_name["Contract Number"]["label"] == "2 Enter the Contract Number"
    # A ticked box reads Selected; its extracted state is kept as it was.
    assert by_name["Negotiated (RFP)"]["choice"] == "Selected"
    # Printed text repeating a box's value is the box read twice: not a second field.
    assert [f["effective_value"] for f in _fields(layout)].count("FA300224C0008") == 1

    rows = [r.values for r in original_rows(layout) if r.kind in ("field", "subfield")]
    assert ["1", "Rating", "DO-C3"] in rows
    assert ["", "Page Number", "1"] in rows
    assert ["2", "Contract Number", "FA300224C0008"] in rows
    assert ["4", "Negotiated (RFP)", "Selected"] in rows
    # No instruction survives as a field name.
    assert not [r for r in rows if r[1].lower().startswith(("enter ", "click "))]


def _ruled_form_pdf() -> bytes:
    """A ruled, fillable form in the shape of an SF33 cover, as the page
    prints it: captions in their own cells, a box whose description gives
    the wrong item number (5 where the page prints 7), a code box drawn
    inside item 7, an unnumbered page box, ticked and unticked choices, a
    table of contents printed in two halves with their own rows, and an
    amendments table whose halves share rows."""

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)

    def line(x0, y0, x1, y1):
        page.draw_line((x0, y0), (x1, y1), width=0.5)

    def text(x, y, words):
        page.insert_text((x, y), words, fontsize=7)

    def box(label, value, rect, kind=fitz.PDF_WIDGET_TYPE_TEXT):
        widget = fitz.Widget()
        widget.field_type = kind
        widget.field_name = f"f{len(list(page.widgets() or []))}"
        widget.field_label = label
        widget.rect = fitz.Rect(*rect)
        widget.field_value = value
        page.add_widget(widget)

    choice = fitz.PDF_WIDGET_TYPE_CHECKBOX
    # Frame and rows.
    line(20, 40, 20, 206)
    line(592, 40, 592, 206)
    for y in (40, 80, 100, 116, 128, 164, 206):
        line(20, y, 592, y)
    # Item 7 (its box says 5) and its CODE box; item 4's two choices.
    line(300, 40, 300, 80)
    text(24, 50, "7. ISSUED BY")
    text(170, 49, "CODE")
    line(200, 40, 200, 52)
    line(200, 52, 296, 52)
    box("Enter the Issued By Code", "FA3002", (204, 42, 294, 51))
    box("5 Enter the info for who issued this solicitation, offer, and award",
        "FA3002  338 ESS CC\nJBSA RANDOLPH, TX 78150-4300\nUNITED STATES", (24, 56, 296, 78))
    text(304, 50, "4. TYPE OF SOLICITATION")
    box('4 Click here to select, "Sealed Bid (IFB)" for the Type of Solicitation', False, (304, 56, 312, 64), choice)
    text(316, 63, "SEALED BID (IFB)")
    box('4 Click here to select, "Negotiated (RFP)" for the Type of Solicitation', True, (304, 68, 312, 76), choice)
    text(316, 75, "NEGOTIATED (RFP)")
    # A page box with its own label and no number.
    text(24, 92, "PAGE")
    box("Enter the page number for this page", "1", (60, 84, 100, 98))
    # Item 11: (X) | SEC | DESCRIPTION | PAGE(S), twice side by side.
    text(250, 112, "11. TABLE OF CONTENTS")
    for x in (45, 70, 250, 300, 325, 350, 530):
        line(x, 116, x, 164)
    for x, words in ((24, "(X)"), (48, "SEC"), (75, "DESCRIPTION"), (255, "PAGE(S)"),
                     (304, "(X)"), (328, "SEC"), (355, "DESCRIPTION"), (535, "PAGE(S)")):
        text(x, 125, words)
    for y in (140, 152):
        line(20, y, 300, y)
    line(300, 146, 592, 146)
    for top, letter, words, ticked, pages in ((128, "A", "SOLICITATION/CONTRACT FORM", True, "1-2"),
                                              (140, "B", "SUPPLIES OR SERVICES", True, "3-10"),
                                              (152, "C", "DESCRIPTION/SPECIFICATIONS", False, "")):
        box(f'11 Click here to select, "{letter} - {words.title()}" for Part I - The Schedule', ticked, (24, top + 2, 32, top + 10), choice)
        text(55, top + 9, letter)
        text(75, top + 9, words)
        box(f"11 Enter the Page numbers for Section {letter}", pages, (255, top + 1, 298, top + 11))
    for top, letter, words, ticked, pages in ((128, "I", "CONTRACT CLAUSES", True, "11-20"),
                                              (146, "J", "LIST OF ATTACHMENTS", False, "")):
        box(f'11 Click here to select, "{letter} - {words.title()}" for Part II', ticked, (304, top + 4, 312, top + 12), choice)
        text(335, top + 11, letter)
        text(355, top + 11, words)
        box(f"11 Enter the Page numbers for Section {letter}", pages, (535, top + 3, 590, top + 13))
    # Item 14: a caption column, then AMENDMENT NUMBER | DATE, twice, on shared rows.
    text(24, 175, "14. ACKNOWLEDGMENT OF AMENDMENTS")
    line(150, 164, 150, 206)
    for x in (260, 370, 480):
        line(x, 164, x, 206)
    for y in (178, 192):
        line(150, y, 592, y)
    for x, words in ((155, "AMENDMENT NUMBER"), (265, "DATE"), (375, "AMENDMENT NUMBER"), (485, "DATE")):
        text(x, 174, words)
    for index, (top, x, number, date) in enumerate(((178, 152, "0009", "08 Feb 2024"), (178, 372, "0007", "16 Jan 2024"),
                                                    (192, 152, "0001", "19 May 2023"), (192, 372, "0002", "05 Jun 2023"))):
        ordinal = ("first", "second", "third", "fourth")[index]
        box(f"14 Enter the {ordinal} Amendment Number (if applicable)", number, (x, top + 2, x + 106, top + 12))
        box(f"14 Enter the acknowledgment date for the {ordinal} Amendment Number", date, (x + 110, top + 2, x + 216, top + 12))
    return doc.tobytes()


def _form_layout(document_id: str, metadata: list[DocumentMetadataField] | None = None, geometry: bool = True) -> dict:
    from app.source_structure.original_layout import _form_pages

    database = SessionLocal()
    try:
        document = database.get(Document, document_id)
        structure = get_or_build_source_structure(database, document)
        return build_original_layout(structure, metadata or [], forms=_form_pages(document, structure) if geometry else None)
    finally:
        database.close()


@pytest.fixture(scope="module")
def ruled_form_id() -> str:
    return _upload(_ruled_form_pdf(), "sf33-ruled")


def test_form_boxes_take_the_item_number_the_page_prints(ruled_form_id: str) -> None:
    by_name = {f["field_name"]: f for f in _fields(_form_layout(ruled_form_id))}
    issued = by_name["Issued By"]
    # Printed as item 7; the box's own description says 5, and that is kept.
    assert issued["item"]["number"] == "7"
    assert issued["form_item"]["number_source"] == "printed" and issued["form_item"]["described_as"] == "5"
    assert issued["label"].startswith("5 Enter the info")
    # The code box drawn inside item 7 is item 7's, though it has no number of its own.
    assert by_name["Issued By Code"]["item"]["id"] == issued["item"]["id"]
    # Both under the printed caption, which names the item.
    assert issued["item"]["name"] == "Issued By" and issued["item"]["has_subfields"]
    # A box with its own label and no printed number stands alone: no parent is invented.
    assert by_name["Page Number"]["item"] is None
    assert by_name["Page Number"]["form_item"]["number_source"] == "none"


def test_without_the_page_geometry_numbers_fall_back_to_the_box_descriptions(ruled_form_id: str) -> None:
    by_name = {f["field_name"]: f for f in _fields(_form_layout(ruled_form_id, geometry=False))}
    # The documented fallback: the description's own number, said to be so.
    assert by_name["Issued By"]["item"]["number"] == "5"
    assert by_name["Issued By"]["form_item"]["number_source"] == "form_field"


def test_ticked_and_unticked_choices_are_both_shown_with_their_raw_state(ruled_form_id: str) -> None:
    by_name = {f["field_name"]: f for f in _fields(_form_layout(ruled_form_id))}
    sealed, negotiated = by_name["Sealed Bid (IFB)"], by_name["Negotiated (RFP)"]
    assert (sealed["choice"], negotiated["choice"]) == ("Not Selected", "Selected")
    # A confirmed state, not a blank or missing value; the form's own value is kept.
    assert sealed["form_state"]["state"] == "not_selected" and sealed["quality"] == "confirmed"
    assert negotiated["form_state"]["raw"] == negotiated["form_state"]["on_state"]
    assert sealed["item"]["name"] == "Type of Solicitation" == negotiated["item"]["name"]

    # The structure reader (and so the Advanced Grid) is unchanged: it still
    # has no candidate for an unticked box.
    database = SessionLocal()
    try:
        structure = get_or_build_source_structure(database, database.get(Document, ruled_form_id))
    finally:
        database.close()
    widgets = [c for c in structure.field_candidates if c.structural_relation == "form_widget"]
    assert widgets and all(c.raw_value.strip() for c in widgets)
    assert not any("Sealed Bid" in c.label_text for c in widgets)


def test_an_unrecognized_choice_state_is_not_read_as_selected_or_not() -> None:
    from app.source_structure.form_geometry import FormBox

    def state(value: str) -> str | None:
        return FormBox("b", "x", "CheckBox", value, "Yes", (0, 0, 1, 1)).state

    assert (state("Yes"), state("Off"), state(""), state("Maybe")) == ("selected", "not_selected", "not_selected", "unrecognized")
    # A box whose own "on" value is "0" (FA30's clause boxes) is ticked when it holds "0".
    assert FormBox("b", "x", "CheckBox", "0", "0", (0, 0, 1, 1)).state == "selected"


def test_a_form_table_keeps_its_printed_columns_and_reading_order(ruled_form_id: str) -> None:
    tables = {t["item"]["number"]: t for t in _tables(_form_layout(ruled_form_id)) if t.get("item")}
    contents = tables["11"]
    assert contents["title"] == "Table of Contents"
    # Printed in two halves with their own rows: one table, left half then right.
    assert contents["headers"] == ["(X)", "SEC", "DESCRIPTION", "PAGE(S)"] and contents["form_table"]["folded"]
    assert [r["cells"] for r in contents["rows"]] == [
        ["Selected", "A", "SOLICITATION/CONTRACT FORM", "1-2"],
        ["Selected", "B", "SUPPLIES OR SERVICES", "3-10"],
        ["Not Selected", "C", "DESCRIPTION/SPECIFICATIONS", ""],
        ["Selected", "I", "CONTRACT CLAUSES", "11-20"],
        ["Not Selected", "J", "LIST OF ATTACHMENTS", ""],
    ]
    # Halves sharing their rows keep the printed columns as they are.
    amendments = tables["14"]
    assert amendments["headers"] == ["AMENDMENT NUMBER", "DATE", "AMENDMENT NUMBER", "DATE"]
    assert not amendments["form_table"]["folded"]
    assert [r["cells"] for r in amendments["rows"]] == [
        ["0009", "08 Feb 2024", "0007", "16 Jan 2024"], ["0001", "19 May 2023", "0002", "05 Jun 2023"],
    ]
    # The table's boxes aren't also listed as fields; each row can be traced to them.
    names = {f["field_name"] for f in _fields(_form_layout(ruled_form_id))}
    assert not any("Page Numbers for Section" in n or "Amendment Number" in n for n in names)
    assert all(r["form_boxes"] and r["bbox"] for r in contents["rows"])
    assert all(not r.get("flags") for t in tables.values() for r in t["rows"])


def test_an_address_block_is_split_into_its_parts_and_kept_whole() -> None:
    from app.source_structure.original_layout import block_parts

    parts = block_parts(
        "FA3002  338 ESS CC\nCP 210-652-6903, 2021 FIRST STREET WEST  BLDG 853\nJBSA RANDOLPH, TX 78150-4300\nUNITED STATES\n"
        "KEVIN SARGENT, Contracting Officer, Email: kevin.sargent.7@us.af.mil Telephone: 487-8228"
    )
    assert [(p["label"], p["value"]) for p in parts] == [
        ("Code", "FA3002"), ("Office", "338 ESS CC"), ("Telephone (CP)", "210-652-6903"),
        ("Street Address", "2021 FIRST STREET WEST BLDG 853"), ("City", "JBSA RANDOLPH"), ("State", "TX"),
        ("ZIP Code", "78150-4300"), ("Country", "UNITED STATES"), ("Contracting Officer", "KEVIN SARGENT"),
        ("Contracting Officer Email", "kevin.sargent.7@us.af.mil"), ("Contracting Officer Telephone", "487-8228"),
    ]
    payment = block_parts("ACCTG DISB STA NR 387700\nDFAS DEAMS, 27 ARKANSAS RD\nLIMESTONE, ME 04751-6216\nUNITED STATES")
    assert ("Office", "DFAS DEAMS") in [(p["label"], p["value"]) for p in payment]
    # Nothing is guessed: no city line, or a line it can't place, and the block stays whole.
    assert block_parts("SCD:    PAS:") == []
    assert block_parts("ACME CORP\nSomewhere in the building\nAUSTIN, TX 78701") == []


def test_extractor_rows_repeating_the_form_are_kept_but_not_shown(ruled_form_id: str) -> None:
    from app.source_structure.original_export import original_rows

    metadata = [
        _metadata(ruled_form_id, "pages", "DESCRIPTION", "3-10 11-20", 1),  # boxes' values
        _metadata(ruled_form_id, "toc_row", "SEC", "SOLICITATION/CONTRACT FORM 1-2", 1),  # a form row read as text
        _metadata(ruled_form_id, "page_label", "PAGE", "", 1),  # a printed label, no value
        _metadata(ruled_form_id, "contractor", "Contractor Name", "ACME FEDERAL LLC", 1),  # a distinct value
    ]
    layout = _form_layout(ruled_form_id, metadata)
    rows = {f["label"]: f for f in _fields(layout) if f.get("placed_by_page")}
    assert rows["DESCRIPTION"]["form_duplicate"]["code"] == "form_value"
    assert rows["SEC"]["form_duplicate"]["code"] == "form_row"
    assert rows["PAGE"]["form_duplicate"]["code"] == "form_label"
    assert "form_duplicate" not in rows["Contractor Name"]
    # The reader's grid over the form's ruled cells is a repeat too.
    grid = [t for t in _tables(layout) if t["detection_method"] == "pdf_ruling_lines"]
    assert all(t["form_duplicate"]["code"] == "form_grid" for t in grid)
    assert layout["stats"]["form_duplicates_hidden"] == 3 + len(grid)
    # The records stay in the layout (review and Smart Review read them);
    # the Original File writes the form's boxes instead.
    written = original_rows(layout)
    field_rows = [r.values for r in written if r.kind in ("field", "subfield")]
    assert ["", "Contractor Name", "ACME FEDERAL LLC"] in field_rows
    assert not [r for r in field_rows if r[1] in ("DESCRIPTION", "SEC", "PAGE")]
    assert not [r for r in written if r.values and str(r.values[0]).startswith("7. ISSUED BY")]


def test_the_form_reads_the_same_in_the_tab_layout_and_the_excel_export(ruled_form_id: str) -> None:
    import io

    from openpyxl import load_workbook

    from app.source_structure.original_export import original_rows

    layout = client.get(f"/api/documents/{ruled_form_id}/original-structure").json()
    response = client.get(f"/api/documents/{ruled_form_id}/original-structure/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Original File"]
    sheet = workbook["Original File"]
    assert not sheet.merged_cells.ranges
    written = [[c if c is not None else "" for c in row] for row in sheet.iter_rows(values_only=True)]
    rows = original_rows(layout)
    # The workbook is exactly the export's rows.
    assert [w[: len(r.values)] for w, r in zip(written, rows)] == [[v if v is not None else "" for v in r.values] for r in rows]
    # Item numbers once, in the order the layout gives them.
    printed_numbers = [r.values[0] for r in rows if r.kind in ("item", "field", "table_title") and r.values[0][:1].isdigit()]
    layout_numbers = []
    for section in layout["sections"]:
        for item in section["items"]:
            number = (item.get("item") or {}).get("number")
            if number and (not layout_numbers or layout_numbers[-1] != number):
                layout_numbers.append(number)
    assert printed_numbers == layout_numbers
    # In the form's numbering (item 4, printed beside item 7, comes first).
    assert printed_numbers == ["4", "7", "11", "14"]
    assert ["7", "Issued By", ""] == written[[r[0] for r in written].index("7")][:3]
    assert ["", "Sealed Bid (IFB)", "Not Selected"] in [r[:3] for r in written]
    assert ["11", "Table of Contents", "Source: page 1"] in [r[:3] for r in written]
    assert ["Not Selected", "C", "DESCRIPTION/SPECIFICATIONS", None] in [list(r[:4]) for r in sheet.iter_rows(values_only=True)]
    # The address block: its parts under it, its printed text kept in a note.
    issued_row = next(i for i, r in enumerate(written) if r[1] == "Issued By" and r[0] == "")
    assert written[issued_row + 1][1:3] == ["Code", "FA3002"]
    assert "Printed as FA3002" in sheet.cell(row=issued_row + 1, column=3).comment.text


def test_original_structure_endpoint(brief_id: str) -> None:
    response = client.get(f"/api/documents/{brief_id}/original-structure")
    assert response.status_code == 200
    body = response.json()
    assert body["stats"]["fields"] >= len(expected_fields())
    assert body["stats"]["tables"] == 3
    assert client.get(f"/api/documents/{uuid4()}/original-structure").status_code == 404


def test_rows_and_unextracted_fields_are_reviewed_without_losing_the_printed_value(brief_id: str) -> None:
    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    field = next(f for f in _fields(layout) if f["metadata"] is None and f["document_value"])
    table = _tables(layout)[0]
    row = table["rows"][1]
    assert field["record_key"].startswith("field:") and row["record_key"].startswith(f"{table['record_key']}:row:")

    corrected = ["X" + c for c in row["cells"]]
    response = client.post(
        f"/api/documents/{brief_id}/original-structure/review",
        json={"records": [row["record_key"]], "action": "correct", "cells": corrected, "reason": "typo"},
    )
    assert response.status_code == 200 and response.json()["reviewed"] == 1
    response = client.post(
        f"/api/documents/{brief_id}/original-structure/review",
        json={"records": [field["record_key"], "field:999:none:missing:0"], "action": "reject"},
    )
    results = {r["record_key"]: r for r in response.json()["results"]}
    assert results[field["record_key"]]["ok"] and not results["field:999:none:missing:0"]["ok"]

    after = client.get(f"/api/documents/{brief_id}/original-structure").json()
    same_row = next(r for t in _tables(after) for r in t["rows"] if r["record_key"] == row["record_key"])
    assert same_row["cells"] == row["cells"]  # the printed row is untouched
    assert same_row["reviewed_cells"] == corrected and same_row["review_status"] == "Corrected"
    same_field = next(f for f in _fields(after) if f["record_key"] == field["record_key"])
    assert same_field["review_status"] == "Rejected"
    assert same_field["structure_review"]["original_value"] == field["document_value"]

    # Approving the corrected row keeps the correction; history keeps both.
    client.post(f"/api/documents/{brief_id}/original-structure/review", json={"records": [row["record_key"]], "action": "accept"})
    final = next(r for t in _tables(client.get(f"/api/documents/{brief_id}/original-structure").json()) for r in t["rows"] if r["record_key"] == row["record_key"])
    assert final["review_status"] == "Approved" and final["reviewed_cells"] == corrected


def test_a_row_correction_needs_every_column(brief_id: str) -> None:
    row = _tables(client.get(f"/api/documents/{brief_id}/original-structure").json())[0]["rows"][0]
    response = client.post(
        f"/api/documents/{brief_id}/original-structure/review",
        json={"records": [row["record_key"]], "action": "correct", "cells": ["only one"]},
    )
    assert response.json()["results"][0]["ok"] is False


def test_fa30_style_contract_parts_open_sections_and_values_never_do() -> None:
    from app.source_structure.models import StructuredRegion, StructureStats, StructuredSourceDocument

    def region(page: int, y: float, kind: str, text: str) -> StructuredRegion:
        return StructuredRegion(region_id=f"r{page}{y}", region_type=kind, text=text, normalized_text=text.lower(), page=page, bbox=(10, y, 200, y + 10), extraction_method="native")

    structure = StructuredSourceDocument(
        document_id="d", source_type="pdf", extractor_version=0, field_candidates=[], table_candidates=[],
        regions=[
            region(1, 60, "HEADING", "F2X3C34150A001"),  # a requisition number set in bold
            region(1, 80, "HEADING", "SOLICITATION"),
            region(2, 40, "HEADING", "Section A - Solicitation/Contract Form"),
            region(3, 40, "PARAGRAPH", "Section B - Supplies or Services & Prices\nAdditional Information"),
            region(3, 90, "HEADING", "(Dec 2018)"),
            region(3, 120, "HEADING", "Unit Price"),
            region(4, 40, "HEADING", "Section C - Description/Specifications"),
            region(4, 80, "HEADING", "Key Position"),
        ],
        stats=StructureStats(pages=4),
    )
    from app.source_structure.original_layout import _section_headings

    headings = [(title, level) for _, _, title, level in _section_headings(structure, [])]
    assert headings == [
        ("SOLICITATION", 2),
        ("Section A - Solicitation/Contract Form", 1),
        ("Section B - Supplies or Services & Prices", 1),
        ("Section C - Description/Specifications", 1),
        ("Key Position", 2),
    ]


def test_selected_records_export_in_every_format(brief_id: str) -> None:
    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    field = _fields(layout)[0]
    row = _tables(layout)[0]["rows"][0]
    for fmt in ("xlsx", "csv", "json"):
        response = client.post(
            f"/api/documents/{brief_id}/original-structure/export-selected",
            json={"records": [field["record_key"], row["record_key"]], "format": fmt},
        )
        assert response.status_code == 200, response.text
    data = response.json()
    assert data["stats"]["fields"] == 1 and data["stats"]["table_rows"] == 1
    csv_text = client.post(
        f"/api/documents/{brief_id}/original-structure/export-selected", json={"records": [field["record_key"]], "format": "csv"},
    ).content.decode("utf-8-sig").splitlines()
    # The same rows as the Excel export: the section, the field columns, the
    # field's item, the field.
    assert csv_text[1] == "Item #,Field Name,Extracted Value" and len(csv_text) == 4


def test_a_missing_row_is_added_marked_and_still_needs_approval(brief_id: str) -> None:
    table = _tables(client.get(f"/api/documents/{brief_id}/original-structure").json())[0]
    printed = [r["cells"] for r in table["rows"] if not r.get("added")]
    cells = ["New"] + [""] * (len(table["headers"]) - 1)
    response = client.post(
        f"/api/documents/{brief_id}/original-structure/review",
        json={"records": [table["record_key"]], "action": "add_row", "cells": cells, "reason": "missed by OCR"},
    ).json()
    assert response["reviewed"] == 1
    key = response["results"][0]["record_key"]
    after = next(t for t in _tables(client.get(f"/api/documents/{brief_id}/original-structure").json()) if t["record_key"] == table["record_key"])
    added = next(r for r in after["rows"] if r["record_key"] == key)
    assert added["added"] and added["cells"] == cells and added["review_status"] == "Added"
    assert [r["cells"] for r in after["rows"] if not r.get("added")] == printed
    client.post(f"/api/documents/{brief_id}/original-structure/review", json={"records": [key], "action": "accept"})
    after = next(t for t in _tables(client.get(f"/api/documents/{brief_id}/original-structure").json()) if t["record_key"] == table["record_key"])
    assert next(r for r in after["rows"] if r["record_key"] == key)["review_status"] == "Approved"
    # Adding needs a whole row; a table itself isn't approved.
    bad = client.post(f"/api/documents/{brief_id}/original-structure/review",
                      json={"records": [table["record_key"]], "action": "accept"}).json()
    assert bad["results"][0]["ok"] is False


def test_record_history_lists_every_decision_oldest_first(brief_id: str) -> None:
    layout = client.get(f"/api/documents/{brief_id}/original-structure").json()
    row = next(r for t in _tables(layout) for r in t["rows"] if not r.get("structure_review") and not r.get("added"))
    for action, extra in (("correct", {"cells": ["Z" + c for c in row["cells"]]}), ("accept", {})):
        client.post(f"/api/documents/{brief_id}/original-structure/review",
                    json={"records": [row["record_key"]], "action": action, "reason": action, **extra})
    history = client.get(f"/api/documents/{brief_id}/original-structure/history", params={"record_key": row["record_key"]}).json()
    assert [e["new_status"] for e in history["entries"]] == ["edited", "accepted"]
    assert history["entries"][-1]["previous_value"] == row["cells"]  # the printed row, kept
    empty = client.get(f"/api/documents/{brief_id}/original-structure/history", params={"record_key": "field:9:x:y:0"}).json()
    assert empty["entries"] == []


def _grid(rows: list[list[str]], *, header: list[str] | None = None, method: str = "pdf_ruling_lines"):
    from app.source_structure.models import TableCandidate, TableCell

    return TableCandidate(
        candidate_id="t", region_id="t", page=1, bbox=(0, 0, 100, 100), detection_method=method,
        headers=header or [f"Column {i + 1}" for i in range(len(rows[0]))],
        header_cells=[TableCell(row_index=-1, column_index=i, text=h) for i, h in enumerate(header or [])],
        rows=[[TableCell(row_index=r, column_index=c, text=t) for c, t in enumerate(row)] for r, row in enumerate(rows)],
    )


def test_text_cut_into_a_grid_is_set_aside_and_real_tables_kept() -> None:
    from app.source_structure.original_layout import _is_broken_prose

    empty = _grid([["", "", ""]] * 6)
    split = _grid([["[ List name, addr", "ess, telephone", "and contact for suppl", "iers surveyed"], ["", "", "", ""]] * 2)
    sparse = _grid([["", "Product Service Code: S2", "", "", "", ""]] + [[""] * 6] * 9)
    clauses = _grid([["52.202-1", "Definitions.", "Jun 2020"], ["52.204-7", "System for Award Management.", "Nov 2024"], ["52.204-13", "SAM Maintenance.", "Oct 2018"]], method="pdf_column_alignment")
    printed = _grid([["", "", ""], ["", "", ""]], header=["Mod Number", "Description", "Date"])
    assert _is_broken_prose([empty]) and _is_broken_prose([split]) and _is_broken_prose([sparse])
    assert not _is_broken_prose([clauses])
    assert not _is_broken_prose([printed])  # a printed header is always a table

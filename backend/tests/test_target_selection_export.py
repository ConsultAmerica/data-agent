"""Export Selected from Source Verification: only the selected fields and
table rows, read from the persisted extraction results."""

from __future__ import annotations

import csv
import io
import json
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from contract_brief_fixture import contract_brief_pdf

client = TestClient(app)


@pytest.fixture(scope="module")
def extracted() -> tuple[str, dict]:
    body = client.post("/api/documents/upload", files={"file": (f"sel-{uuid4().hex[:6]}.pdf", contract_brief_pdf(), "application/pdf")}).json()
    document_id = body["document_id"]
    job_id = body.get("processing_job_id") or client.post(f"/api/documents/{document_id}/jobs/process").json()["id"]
    client.get(f"/v1/jobs/{job_id}")
    targets = client.get(f"/api/documents/{document_id}/targets").json()["targets"]
    ids = [t["id"] for t in targets if t.get("source_examples") and t["target_type"] not in {"clause", "section", "obligation"}]
    client.get(f"/v1/jobs/{client.post(f'/api/documents/{document_id}/jobs/extract', json={'target_ids': ids}).json()['id']}")
    result = client.get(f"/api/documents/{document_id}/extract-results").json()
    assert result["scalars"] and result["tables"]
    return document_id, result


def _export(document_id: str, **selection):
    return client.post(f"/api/documents/{document_id}/extract-results/export-selected", json=selection)


def test_only_selected_fields_and_rows_are_exported(extracted) -> None:
    document_id, result = extracted
    chosen = [s for s in result["scalars"] if s["value"]][:2]
    table = result["tables"][0]
    response = _export(
        document_id, format="json",
        fields=[s["normalized_key"] for s in chosen],
        rows=[{"table": table["target"], "index": i} for i in (1, 0)],  # ticked out of order
    )
    assert response.status_code == 200, response.text
    body = json.loads(response.content)
    assert [f["key"] for f in body["fields"]] == [s["normalized_key"] for s in chosen]
    assert [f["value"] for f in body["fields"]] == [str(s["value"]) for s in chosen]
    [exported] = body["tables"]
    assert [r["index"] for r in exported["rows"]] == [0, 1]  # stored order
    assert exported["columns"] == table["columns"]
    assert len(result["scalars"]) > len(body["fields"])  # unselected left out


def test_xlsx_and_csv_hold_the_same_selection(extracted) -> None:
    document_id, result = extracted
    key = next(s["normalized_key"] for s in result["scalars"] if s["value"])
    table = result["tables"][0]
    xlsx = _export(document_id, format="xlsx", fields=[key], rows=[{"table": table["target"], "index": 0}])
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book["Selected fields"].max_row == 2  # header + one field
    assert book["Table 1"].max_row == 3  # caption, headers, one row
    text = _export(document_id, format="csv", fields=[key], rows=[]).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert len(rows) == 2 and rows[1][0] == "Field"


def test_unknown_or_empty_selections_are_refused(extracted) -> None:
    document_id, _ = extracted
    assert _export(document_id, format="csv", fields=["no_such_field"]).status_code == 422
    assert _export(document_id, format="csv", rows=[{"table": "no_such_table", "index": 0}]).status_code == 422
    assert _export(document_id, format="csv").status_code == 422

"""Custom export repair: a 500-line invoice through the real pipeline,
preview = export, formula-injection protection, transforms that never
invent values."""

from __future__ import annotations

import csv
import io
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database.session import SessionLocal
from app.main import app
from app.models.document import Document
from app.staging.custom_export import _transform, safe_cell
from app.staging.preparation import prepare_staging
from tests.invoice_fixtures import long_invoice_pdf

client = TestClient(app)


@pytest.fixture(scope="module")
def invoice() -> dict:
    marker = uuid4().hex[:8]
    upload = client.post("/api/documents/upload", files={"file": (f"long-invoice-{marker}.pdf", long_invoice_pdf(marker=marker), "application/pdf")})
    assert upload.status_code == 201, upload.text
    document_id = upload.json()["document_id"]
    assert client.post(f"/api/documents/{document_id}/extract-pages", json={"run_ocr": False}).status_code == 200
    database = SessionLocal()
    try:
        prepare_staging(database, database.get(Document, document_id))
    finally:
        database.close()
    workbook = client.get(f"/api/documents/{document_id}/staging-workbook").json()
    summary = next(d for d in workbook["datasets"] if d["dataset_id"] == "invoice_summary")
    number_field = next(c["canonical_field"] for c in summary["columns"] if c["canonical_field"].endswith("invoice_number"))
    return {"id": document_id, "summary_record": summary["records"][0]["record_id"], "number_field": number_field}


def manifest(invoice: dict, fmt: str = "csv", **column_extra) -> dict:
    line = lambda field: {"dataset_id": "invoice_lines", "field": f"invoice.line.{field}"}  # noqa: E731
    return {
        "format": fmt,
        "include_header": True,
        "columns": [
            {"label": "Invoice", "sources": [{"dataset_id": "invoice_summary", "field": invoice["number_field"], "record_id": invoice["summary_record"]}]},
            {"label": "Item", "sources": [line("item_number")]},
            {"label": "Description", "sources": [line("description")], **column_extra},
            {"label": "Qty x Price", "sources": [line("quantity"), line("unit_price")], "separator": " x "},
            {"label": "Amount", "sources": [line("amount")]},
            {"label": "PO line", "sources": [line("po_line")]},
        ],
    }


def test_all_500_lines_export_once_each_with_the_invoice_number_as_a_column(invoice) -> None:
    response = client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom", json=manifest(invoice))
    assert response.status_code == 200, response.text
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == ["Invoice", "Item", "Description", "Qty x Price", "Amount", "PO line"]
    body = rows[1:]
    assert len(body) == 500  # one row per line — no header rows per line, none lost
    assert {row[0] for row in body} == {"INV-500"}
    assert [row[1] for row in body[:3]] == ["0001", "0002", "0003"]  # leading zeros kept
    assert body[0][3] == "2 x 11.25" and body[0][4] == "22.50"  # combined; decimals kept
    assert body[-1][1] == "0500"
    assert all(row[5] == "" for row in body)  # missing stays empty, never filled in


def test_xlsx_has_the_same_500_rows_as_text_where_extracted_as_text(invoice) -> None:
    response = client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom", json=manifest(invoice, "xlsx"))
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 501
    assert sheet["B2"].value == "0001" and sheet["B501"].value == "0500"


def test_the_preview_is_the_export(invoice) -> None:
    body = manifest(invoice)
    preview = client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom/preview", json=body).json()
    exported = list(csv.reader(io.StringIO(client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom", json=body).content.decode("utf-8-sig"))))
    assert preview["total_rows"] == 500 and preview["row_dataset"] == "invoice_lines"
    assert preview["headers"] == exported[0]
    assert preview["rows"] == exported[1:26]


def test_transforms_change_text_but_never_invent_values(invoice) -> None:
    upper = client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom/preview", json=manifest(invoice, transform="upper")).json()
    assert upper["rows"][0][2].startswith("PART 0001")
    assert _transform(None, "upper") is None and _transform("", "date_iso") == ""
    assert _transform("October 2, 2026", "date_iso") == "2026-10-02"
    assert _transform("03/04/2026", "date_iso") == "03/04/2026"  # ambiguous: as printed
    assert _transform("not a date", "date_iso") == "not a date"
    assert _transform("  padded  ", "trim") == "padded"


def test_formula_like_values_and_headers_are_written_as_text(invoice) -> None:
    assert safe_cell("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    assert safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell("+1+1") == "'+1+1"
    assert safe_cell("-12.50") == "-12.50" and safe_cell("0001") == "0001" and safe_cell(42) == 42
    body = manifest(invoice)
    body["columns"][1]["label"] = "=cmd|' /C calc'!A0"
    rows = list(csv.reader(io.StringIO(client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom", json=body).content.decode("utf-8-sig"))))
    assert rows[0][1] == "'=cmd|' /C calc'!A0"
    sheet = load_workbook(io.BytesIO(client.post(f"/api/documents/{invoice['id']}/staging-workbook/export-custom", json={**body, "format": "xlsx"}).content)).active
    assert sheet["B1"].value == "'=cmd|' /C calc'!A0" and sheet["B1"].data_type == "s"


def test_multiline_and_unicode_values_survive(invoice) -> None:
    from app.staging.custom_export import _csv, _xlsx

    text = _csv(["Note"], [["Línea 1\nZürich “Ω” 0007"]], True)
    assert list(csv.reader(io.StringIO(text)))[1] == ["Línea 1\nZürich “Ω” 0007"]
    sheet = load_workbook(io.BytesIO(_xlsx(["Note"], [["Línea 1\nZürich “Ω” 0007"]], True))).active
    assert sheet["A2"].value == "Línea 1\nZürich “Ω” 0007"
